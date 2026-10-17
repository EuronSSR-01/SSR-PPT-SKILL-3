#!/usr/bin/env python3
"""Render decks to images for visual review.

  preview.py html  deck.html  --out DIR   per-slide PNGs (fully revealed), console/network/offline checks,
                                          keyboard test, small-viewport shot, print PDF, contact sheets
  preview.py pptx  deck.pptx  --out DIR   LibreOffice -> PDF -> PNGs (or --export wps|powerpoint to have that
                                          app export the PDF, or --pdf FILE exported from the deck), layout
                                          check, native object inventory, fonts, optional --edit-test on a
                                          copy, contact sheets
  preview.py sheet PNGS|DIR   --out FILE  contact sheet plus -gray and -blur variants
  preview.py compare A B ...  --out FILE  images side by side for comparing directions

html and pptx also report body text per slide (and how much of it is sentences rather than short labels);
with --visual-first they point at long paragraphs, boxes of points and slides with many sentences, whose
information should move into pictures, not be deleted (hints, not a pass/fail check). --zoom also writes four enlarged quarters of every slide for a second, closer look.
The layout check flags text past its box, one-character lines and a content frame that sits off-centre on
three or more slides (warnings to confirm by looking).

Dependencies: Pillow; html needs playwright (bundled Chromium or installed Chrome/Edge);
pptx needs PyMuPDF and LibreOffice (or --pdf), and python-pptx for the inventory and --edit-test.
Every image this writes must still be opened and looked at; numbers in report.json are only hints.
"""
import argparse
import hashlib
import importlib
from io import BytesIO
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def need(module, hint):
    try:
        return importlib.import_module(module)
    except ImportError:
        sys.exit(f"missing dependency {module!r}: {hint}")


def pymupdf_or_none():
    for name in ("pymupdf", "fitz"):  # newer releases deprecate the fitz name
        try:
            return importlib.import_module(name)
        except ImportError:
            pass
    return None


# ---------------------------------------------------------------- contact sheets

def _font(size):
    from PIL import ImageFont
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def make_sheet(images, out, cols=None, thumb_w=400, labels=None):
    """Grid of thumbnails, plus -gray and -blur variants for hierarchy checks."""
    Image = need("PIL.Image", "python -m pip install pillow")
    from PIL import ImageDraw, ImageFilter
    images = [Path(p) for p in images]
    if not images:
        raise SystemExit("no images for contact sheet")
    cols = cols or min(4, max(1, math.ceil(math.sqrt(len(images) * 1.3))))
    thumbs = []
    for p in images:
        im = Image.open(p).convert("RGB")
        im.thumbnail((thumb_w, thumb_w * 4))
        thumbs.append(im)
    cell_h = max(t.height for t in thumbs)
    pad, label_h = 16, 22
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new("RGB", (pad + cols * (thumb_w + pad), pad + rows * (cell_h + label_h + pad)), (38, 38, 40))
    draw = ImageDraw.Draw(sheet)
    font = _font(15)
    for i, t in enumerate(thumbs):
        x = pad + (i % cols) * (thumb_w + pad)
        y = pad + (i // cols) * (cell_h + label_h + pad)
        draw.text((x, y), labels[i] if labels else f"{i + 1:02d}", fill=(220, 220, 220), font=font)
        sheet.paste(t, (x, y + label_h))
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(out)
    sheet.convert("L").save(out.with_name(out.stem + "-gray" + out.suffix))
    sheet.filter(ImageFilter.GaussianBlur(max(2, thumb_w // 45))).save(out.with_name(out.stem + "-blur" + out.suffix))
    return out


def make_compare(images, out, labels=None, height=540):
    Image = need("PIL.Image", "python -m pip install pillow")
    from PIL import ImageDraw
    ims = []
    for p in images:
        im = Image.open(p).convert("RGB")
        ims.append(im.resize((round(im.width * height / im.height), height)))
    pad, label_h = 20, 28
    canvas = Image.new("RGB", (pad + sum(i.width + pad for i in ims), height + label_h + 2 * pad), (38, 38, 40))
    draw = ImageDraw.Draw(canvas)
    font = _font(20)
    x = pad
    for n, im in enumerate(ims):
        draw.text((x, pad), labels[n] if labels else chr(65 + n), fill=(230, 230, 230), font=font)
        canvas.paste(im, (x, pad + label_h))
        x += im.width + pad
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out)
    return out


# ---------------------------------------------------------------- HTML

CAPTURE_CSS = """
.ssr-ui { display: none !important; }
*, *::before, *::after { animation-duration: 0s !important; animation-delay: 0s !important;
  transition-duration: 0s !important; transition-delay: 0s !important; caret-color: transparent !important; }
"""

WAIT_ASSETS_JS = """async () => {
  await document.fonts.ready;
  const svgImages = Array.from(document.querySelectorAll('svg image')).map(el => {
    const img = new Image(); img.src = el.href.baseVal; return img;
  });
  const images = await Promise.all([...document.images, ...svgImages].map(async (img, index) => {
    let error = null;
    try { await img.decode(); } catch (e) { error = String(e); }
    return {index, source: img.src.startsWith('data:') ? 'inline image' : img.src,
            ok: !error && img.naturalWidth > 0 && img.naturalHeight > 0, error};
  }));
  return {fonts_status: document.fonts.status, images,
          fonts: Array.from(document.fonts).map(f => ({family: f.family, status: f.status}))};
}"""


# Geometry checks on the fully revealed slide. They catch common defects without looking at the image
# and tell agents with vision where to look. Opt-outs: data-bleed (intentional bleed off the slide),
# data-lint-ignore (skip a subtree). Coordinates are canvas px relative to the slide.
LAYOUT_JS = r"""(index) => {
  const slide = document.querySelectorAll('#ssr-deck > .ssr-slide')[index];
  const sr = slide.getBoundingClientRect();
  const k = sr.width / (slide.offsetWidth || sr.width) || 1;
  const rel = r => ({x: Math.round((r.left - sr.left) / k), y: Math.round((r.top - sr.top) / k),
                     w: Math.round(r.width / k), h: Math.round(r.height / k)});
  const out = [];
  const label = el => (el.textContent || '').trim().replace(/\s+/g, ' ').slice(0, 40);
  const add = (kind, severity, el, rect, detail) => out.push({kind, severity, text: label(el), rect: rel(rect), detail});
  const skip = el => el.closest('[data-lint-ignore], .ssr-notes, .ssr-ui, script, style');
  const shown = el => { for (let e = el; e && e !== slide; e = e.parentElement) { const cs = getComputedStyle(e);
      if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) < 0.05) return false; } return true; };
  const ownRects = el => { const rects = [];
    for (const n of el.childNodes) if (n.nodeType === 3 && n.textContent.trim()) {
      const range = document.createRange(); range.selectNodeContents(n);
      for (const r of range.getClientRects()) if (r.width > 0.5 && r.height > 0.5) rects.push(r); }
    return rects; };
  const fontPx = el => { const fs = parseFloat(getComputedStyle(el).fontSize) || 16;
    if (el instanceof SVGElement && el.getScreenCTM) { const m = el.getScreenCTM(); if (m) return fs * Math.hypot(m.a, m.b) / k; }
    return fs; };
  const texts = Array.from(slide.querySelectorAll('*')).filter(el => !skip(el) && shown(el) && ownRects(el).length)
    .map(el => ({el, rects: ownRects(el), fs: fontPx(el)}));

  for (const t of texts) {
    // 1. text outside the slide (the slide clips it)
    if (!t.el.closest('[data-bleed]')) for (const r of t.rects)
      if (r.left < sr.left - 1 || r.top < sr.top - 1 || r.right > sr.right + 1 || r.bottom > sr.bottom + 1) {
        add('text_outside_slide', 'error', t.el, r, 'Text runs past the slide edge; add data-bleed if intended.'); break; }
    // 2. text clipped by an overflow-hidden ancestor
    for (let a = t.el; a && a !== slide; a = a.parentElement) {  // the element itself may clip its own text
      const cs = getComputedStyle(a);
      if (cs.overflowX === 'visible' && cs.overflowY === 'visible') continue;
      const ar = a.getBoundingClientRect();
      const hit = t.rects.find(r => r.right > ar.right + 1 || r.bottom > ar.bottom + 1 || r.left < ar.left - 1 || r.top < ar.top - 1);
      if (hit) { add('text_clipped', 'error', t.el, hit, 'Text exceeds an overflow-hidden container.'); break; }
    }
    // 5. font size
    if (t.fs < 12) add('font_too_small', 'warning', t.el, t.rects[0], `Effective size ${t.fs.toFixed(1)}px (< 12px).`);
  }

  // 3. text overlapping other text (em-box height, so tight line-height does not count)
  const box = (r, fs) => { const inset = Math.max(0, (r.height - fs * k) / 2);
    return {l: r.left, r: r.right, t: r.top + inset, b: r.bottom - inset}; };
  const symbolsOnly = el => /^[\p{P}\p{S}\s]+$/u.test(el.textContent || '');  // big quote marks: tiny ink, huge em box
  for (let i = 0; i < texts.length; i++) for (let j = i + 1; j < texts.length; j++) {
    const a = texts[i], b = texts[j];
    if (a.el.contains(b.el) || b.el.contains(a.el) || symbolsOnly(a.el) || symbolsOnly(b.el)) continue;
    let found = null;
    for (const ra of a.rects) { const A = box(ra, a.fs);
      for (const rb of b.rects) { const B = box(rb, b.fs);
        const ix = Math.min(A.r, B.r) - Math.max(A.l, B.l), iy = Math.min(A.b, B.b) - Math.max(A.t, B.t);
        if (ix > 4 && iy > 4) { found = ra; break; } }
      if (found) break; }
    if (found) add('text_overlap', 'error', a.el, found, `Overlaps “${label(b.el)}”.`);
  }

  // 4. orphan line: a wrapped line holding about one character (a trailing CJK punctuation mark aside),
  // whether it ends the block or sits before a forced break
  const blocks = new Set(texts.map(t => t.el instanceof SVGElement ? null :
    t.el.closest('p, li, h1, h2, h3, h4, h5, h6, td, th, dd, blockquote, figcaption, div')).filter(Boolean));
  for (const el of blocks) {
    if (skip(el) || el.querySelector('p, li, h1, h2, h3, h4, h5, h6, td, div, table, svg')) continue;
    const range = document.createRange(); range.selectNodeContents(el);
    const rs = Array.from(range.getClientRects()).filter(r => r.width > 0.5 && r.height > 0.5);
    if (rs.length < 2) continue;
    const fs = fontPx(el) * k, lines = [];
    for (const r of rs) { const c = (r.top + r.bottom) / 2;
      const line = lines.find(l => Math.abs(l.c - c) < fs * 0.5);
      if (line) { line.l = Math.min(line.l, r.left); line.r = Math.max(line.r, r.right); }
      else lines.push({c, l: r.left, r: r.right, t: r.top, b: r.bottom}); }
    if (lines.length < 2) continue;
    lines.sort((a, b) => a.c - b.c);
    const endsInMark = /[。，、；：！？）」』”’》】]\s*$/.test(el.textContent || '');
    for (let i = 1; i < lines.length; i++) {
      const line = lines[i], w = line.r - line.l;
      const last = i === lines.length - 1;
      if (w <= fs * 1.6 || (last && endsInMark && w <= fs * 2.2)) {
        add('orphan_line', 'warning', el, new DOMRect(line.l, line.t, w, line.b - line.t),
            'A line holds about one character; rewrite, widen or break the line deliberately.');
        break; }
    }
  }

  // 6. contrast against whatever is painted behind the text
  const parse = c => { const m = /rgba?\(([^)]+)\)/.exec(c || ''); if (!m) return null;
    const p = m[1].split(/[ ,/]+/).filter(Boolean).map(Number); return {r: p[0], g: p[1], b: p[2], a: p.length > 3 ? p[3] : 1}; };
  const lum = c => { const f = v => { v /= 255; return v <= 0.03928 ? v / 12.92 : Math.pow((v + 0.055) / 1.055, 2.4); };
    return 0.2126 * f(c.r) + 0.7152 * f(c.g) + 0.0722 * f(c.b); };
  const behind = (el, r) => {
    if (!(el instanceof SVGElement)) {  // text paints over its own background first (badges, tags)
      const own = getComputedStyle(el);
      if (own.backgroundImage && own.backgroundImage !== 'none') return null;
      const bg = parse(own.backgroundColor); if (bg && bg.a > 0.9) return bg; }
    const stack = document.elementsFromPoint((r.left + r.right) / 2, (r.top + r.bottom) / 2);
    for (const e of stack) {
      if (e === el || el.contains(e) || !slide.contains(e) && e !== slide) continue;
      if (/^(IMG|VIDEO|CANVAS|IMAGE)$/i.test(e.tagName)) return null;
      const cs = getComputedStyle(e);
      if (e instanceof SVGElement) {
        if (e instanceof SVGGeometryElement) { const f = parse(cs.fill);
          if (f && f.a > 0.9 && parseFloat(cs.fillOpacity || 1) > 0.9) return f;
          if (cs.fill && cs.fill !== 'none' && !f) return null; }
        continue; }
      if (cs.backgroundImage && cs.backgroundImage !== 'none') return null;
      const bg = parse(cs.backgroundColor); if (bg && bg.a > 0.9) return bg;
      if (e === slide) break;
    }
    return null; };
  const lowContrast = new Map();  // one finding per colour pair per slide
  for (const t of texts) {
    const cs = getComputedStyle(t.el);
    const fg = parse(t.el instanceof SVGElement ? cs.fill : cs.color);
    const bg = behind(t.el, t.rects[0]);
    if (!fg || !bg) continue;
    const mix = {r: fg.r * fg.a + bg.r * (1 - fg.a), g: fg.g * fg.a + bg.g * (1 - fg.a), b: fg.b * fg.a + bg.b * (1 - fg.a)};
    const [hi, lo] = [lum(mix), lum(bg)].sort((a, b) => b - a);
    const ratio = (hi + 0.05) / (lo + 0.05);
    const large = t.fs >= 24 || (t.fs >= 18.66 && parseInt(cs.fontWeight, 10) >= 700);
    if (ratio >= (large ? 3 : 4.5)) continue;
    const key = `${Math.round(mix.r)},${Math.round(mix.g)},${Math.round(mix.b)}/${bg.r},${bg.g},${bg.b}/${large}`;
    const seen = lowContrast.get(key);
    if (seen) { seen.count++; continue; }
    lowContrast.set(key, {t, ratio, need: large ? 3 : 4.5, count: 1});
  }
  for (const {t, ratio, need, count} of lowContrast.values())
    add('low_contrast', 'warning', t.el, t.rects[0],
        `Contrast ${ratio.toFixed(2)}:1 (needs ${need}:1)` + (count > 1 ? `; ${count} text items use this colour pair.` : '.'));

  // 7. leftover placeholders
  const placeholder = /lorem ipsum|\bTODO\b|\bTBD\b|X{3,}|占位|待补充|待填写|此处填写|示例文字/i;
  for (const t of texts) if (placeholder.test(t.el.textContent))
    add('placeholder', 'warning', t.el, t.rects[0], 'Looks like placeholder text.');

  // 8. small body text: own text of each element with its size, position in canvas px (judged in Python)
  const own = el => Array.from(el.childNodes).filter(n => n.nodeType === 3).map(n => n.textContent).join(' ');
  const sized = texts.map(t => [own(t.el), t.fs, (t.rects[0].top - sr.top) / k,
                                (t.rects[t.rects.length - 1].bottom - sr.top) / k, rel(t.rects[0])]);
  return {findings: out, sized, height: sr.height / k};
}"""

# Body text per slide for the visual-first hints: leaf text blocks (SVG text lines included) and the card,
# list or SVG group each sits in. Left out: h1/h2, headline-size text (>= 36px), tables, notes and the
# footer band (bottom 12%).
TEXT_JS = r"""(index) => {
  const slide = document.querySelectorAll('#ssr-deck > .ssr-slide')[index];
  const sr = slide.getBoundingClientRect();
  const k = sr.width / (slide.offsetWidth || sr.width) || 1;
  const rel = r => ({x: Math.round((r.left - sr.left) / k), y: Math.round((r.top - sr.top) / k),
                     w: Math.round(r.width / k), h: Math.round(r.height / k)});
  const skip = el => el.closest('[data-lint-ignore], .ssr-notes, .ssr-ui, script, style, table, h1, h2');
  const shown = el => { for (let e = el; e && e !== slide; e = e.parentElement) { const cs = getComputedStyle(e);
      if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) < 0.05) return false; } return true; };
  const fontPx = el => { const fs = parseFloat(getComputedStyle(el).fontSize) || 16;
    if (el instanceof SVGElement && el.getScreenCTM) { const m = el.getScreenCTM(); if (m) return fs * Math.hypot(m.a, m.b) / k; }
    return fs; };
  const leafSel = 'p, li, dd, dt, blockquote, figcaption, h3, h4, h5, h6, div';
  const isCard = e => { if (e instanceof SVGElement) return e.tagName === 'g';
    if (/^(UL|OL)$/.test(e.tagName)) return true; const cs = getComputedStyle(e);
    const bg = /rgba?\(([^)]+)\)/.exec(cs.backgroundColor); const alpha = bg ? (bg[1].split(/[ ,/]+/)[3] ?? 1) : 0;
    return Number(alpha) > 0.05 || parseFloat(cs.borderTopWidth) > 0 || parseFloat(cs.borderLeftWidth) > 0 ||
           (cs.boxShadow && cs.boxShadow !== 'none'); };
  const candidates = Array.from(slide.querySelectorAll(leafSel)).filter(el => !el.closest('svg') && !el.querySelector(leafSel))
    .concat(Array.from(slide.querySelectorAll('svg text')));
  const leaves = [], groups = new Map();
  for (const el of candidates) {
    const text = (el.textContent || '').trim().replace(/\s+/g, ' ');
    if (!text || skip(el) || !shown(el) || fontPx(el) >= 36) continue;
    const r = el.getBoundingClientRect();
    if (r.top - sr.top >= sr.height * 0.88) continue;
    let group = -1;
    for (let a = el.parentElement; a && a !== slide; a = a.parentElement) {
      const ar = a.getBoundingClientRect();
      if (ar.width * ar.height > sr.width * sr.height * 0.4) break;
      if (isCard(a)) { if (!groups.has(a)) groups.set(a, groups.size); group = groups.get(a); break; }
    }
    leaves.push({text, rect: rel(r), group});
  }
  // pictures, illustrations, charts and tables: images, canvases and SVGs big enough not to be an icon
  // (0.4% of the slide), tables of any size; not in the header or footer band
  const visuals = Array.from(slide.querySelectorAll('img, svg, canvas, video, table')).filter(el => {
    if (el.closest('.ssr-notes, .ssr-ui') || !shown(el) || (el.tagName.toLowerCase() === 'svg' && el.parentElement.closest('svg'))) return false;
    const r = el.getBoundingClientRect(), middle = (r.top + r.bottom) / 2 - sr.top;
    if (middle < sr.height * 0.115 || middle > sr.height * 0.92) return false;
    return el.tagName.toLowerCase() === 'table' || r.width * r.height >= 0.004 * sr.width * sr.height;
  }).length;
  return {leaves, groups: Array.from(groups.keys()).map(a => rel(a.getBoundingClientRect())), visuals};
}"""

CJK_CHAR = re.compile(r"[㐀-鿿豈-﫿぀-ヿ가-힯]")
LATIN_WORD = re.compile(r"[A-Za-z0-9]+(?:[.,'’%-][A-Za-z0-9]+)*")
LONG_TEXT = 60  # 字 in one paragraph or one box of points; a Latin word or number counts as two
PAGE_TEXT = 150  # 字 of sentences on one slide
LABEL_TEXT = 20  # a line up to this long is a label (a name, a number, a short caption beside a picture)
SMALL_PX = 16  # body text under this size on a 1280-wide canvas (12pt on a 13.33-inch slide) is small
SMALL_MIN, SMALL_SHARE = 40, 0.3  # warn when at least this many 字 and this share of the body text are small
BARE_TEXT = 100  # 字 of body text on a slide with no picture, illustration, chart or table (visual-first hint)


# Left and right edges of the content between the header band (top 11.5%) and the footer band (bottom 8%):
# text, pictures, SVG and boxes with a fill or border, leaving out full-width bands. Canvas px.
FRAME_JS = r"""(index) => {
  const slide = document.querySelectorAll('#ssr-deck > .ssr-slide')[index];
  const sr = slide.getBoundingClientRect();
  const k = sr.width / (slide.offsetWidth || sr.width) || 1;
  const shown = el => { for (let e = el; e && e !== slide; e = e.parentElement) { const cs = getComputedStyle(e);
      if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) < 0.05) return false; } return true; };
  let l = Infinity, r = -Infinity;
  const take = rect => {
    if (rect.width < 1 || rect.height < 1 || rect.width > sr.width * 0.95) return;
    if (rect.top - sr.top < sr.height * 0.115 || rect.bottom - sr.top > sr.height * 0.92) return;
    l = Math.min(l, rect.left - sr.left); r = Math.max(r, rect.right - sr.left); };
  for (const el of slide.querySelectorAll('*')) {
    if (el.closest('.ssr-notes, .ssr-ui, script, style, [data-lint-ignore], [data-bleed]') || !shown(el)) continue;
    if (/^(img|svg|canvas|video)$/i.test(el.tagName)) { take(el.getBoundingClientRect()); continue; }
    if (el.closest('svg')) continue;
    const cs = getComputedStyle(el);
    const bg = /rgba?\(([^)]+)\)/.exec(cs.backgroundColor);
    const alpha = bg ? Number(bg[1].split(/[ ,/]+/)[3] ?? 1) : 0;
    if (alpha > 0.05 || parseFloat(cs.borderTopWidth) > 0 || parseFloat(cs.borderLeftWidth) > 0 ||
        (cs.backgroundImage && cs.backgroundImage !== 'none')) take(el.getBoundingClientRect());
    for (const n of el.childNodes) if (n.nodeType === 3 && n.textContent.trim()) {
      const range = document.createRange(); range.selectNodeContents(n);
      for (const rect of range.getClientRects()) take(rect); }
  }
  if (!isFinite(l)) return null;
  return {left: l / k, right: (sr.width - r) / k, span: (r - l) / k, width: sr.width / k};
}"""


def deck_frame(frames, unit, tolerance, threshold):
    """One warning when three or more slides share a content frame (the largest such group) whose left and
    right margins differ by more than threshold. frames: [(slide, {left, right, span, width}) or None].
    Slides whose content spans less than 60% of the width (covers, half-page pictures) are left out.
    Asymmetric frames can be deliberate, e.g. a template with a side bar."""
    wide = [(n, f) for n, f in frames if f and f["span"] > 0.6 * f["width"]]
    best, members = None, []
    for n, f in wide:
        group = [m for m, g in wide if abs(g["left"] - f["left"]) <= tolerance and abs(g["right"] - f["right"]) <= tolerance]
        if len(group) > len(members):
            best, members = f, group
    if not best or len(members) < 3 or abs(best["left"] - best["right"]) <= threshold:
        return None
    pages = ", ".join(str(m) for m in members)
    return {"kind": "content_frame_off_centre", "severity": "warning", "slides": members,
            "left": round(best["left"], 1), "right": round(best["right"], 1), "unit": unit,
            "detail": f"Slides {pages} share a content frame with left margin {best['left']:.0f}{unit} and right "
                      f"margin {best['right']:.0f}{unit}. Unless the asymmetry is deliberate, take the content area "
                      "from the template or align both edges to the header (title, logo)."}


def zoom_quarters(image, out_dir, stem):
    """Four overlapping quarters of a slide image (top-left, top-right, bottom-left, bottom-right)."""
    Image = need("PIL.Image", "python -m pip install pillow")
    im = Image.open(image)
    w, h = im.size
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, (x0, y0) in enumerate(((0, 0), (0.45, 0), (0, 0.45), (0.45, 0.45)), 1):
        box = (round(x0 * w), round(y0 * h), round((x0 + 0.55) * w), round((y0 + 0.55) * h))
        path = out_dir / f"{stem}-{i}.png"
        im.crop(box).save(path)
        paths.append(path)
    return paths


def text_units(text):
    """Reading length in 字: one per CJK character, two per Latin word or number."""
    return len(CJK_CHAR.findall(text)) + 2 * len(LATIN_WORD.findall(text))


def text_density(leaves, groups, visual_first):
    """Summarise one slide's body text. leaves: [{text, rect, groups}] where groups lists the enclosing
    blocks (innermost first); groups: [rect]. units counts all body text; sentence_units only lines longer
    than LABEL_TEXT, so short labels beside pictures do not make a slide look text-heavy. With visual_first,
    long paragraphs, text-heavy boxes and slides with many sentences become hints: draw the information
    they hold (illustration, icon with a short label, dot chart, arrow, chart), then cut repeats and move
    explanations to the notes. Information is not deleted. They are not a pass/fail check."""
    for leaf in leaves:
        leaf["units"] = text_units(leaf["text"])
    findings, flagged = [], set()
    for i, leaf in enumerate(leaves):
        if visual_first and leaf["units"] > LONG_TEXT:
            flagged.add(i)
            findings.append({"kind": "long_paragraph", "severity": "hint", "text": leaf["text"][:40],
                             "rect": leaf["rect"], "units": leaf["units"],
                             "detail": f"About {leaf['units']} 字 in one paragraph. Split it into its items and "
                                       "draw what can be drawn (illustration, icon with a short label, dot chart, "
                                       "arrow, chart); move only explanation and repeats to the notes; keep every "
                                       "item and the conditions."})
    members = {}
    for i, leaf in enumerate(leaves):
        for g in leaf.get("groups", []):
            members.setdefault(g, []).append(i)
    for g in sorted(members, key=lambda g: len(members[g])):  # inner boxes first
        idx = members[g]
        units = sum(leaves[i]["units"] for i in idx)
        if visual_first and len(idx) >= 2 and units > LONG_TEXT and not flagged.intersection(idx):
            flagged.update(idx)
            findings.append({"kind": "long_text_block", "severity": "hint", "text": leaves[idx[0]]["text"][:40],
                             "rect": groups[g], "units": units,
                             "detail": f"About {units} 字 of points in one box. Give each point a picture or an "
                                       "icon with a short label, or draw them as one diagram; keep every point."})
    total = sum(l["units"] for l in leaves)
    sentences = sum(l["units"] for l in leaves if l["units"] > LABEL_TEXT)
    if visual_first and sentences > PAGE_TEXT:
        findings.append({"kind": "text_heavy_slide", "severity": "hint", "text": "(whole slide)", "rect": None,
                         "units": sentences,
                         "detail": f"About {sentences} 字 in sentences. Keep every information item on the slide "
                                   "but carry most of them with pictures (illustration, icon with a short label, "
                                   "dot chart, arrow, chart); move only explanations and repeats to the notes. "
                                   "Short labels are fine and not counted."})
    return {"units": total, "sentence_units": sentences,
            "longest_paragraph": max((l["units"] for l in leaves), default=0), "findings": findings}


def bare_slide(summary, visuals, visual_first):
    """Record a slide's pictures (photos, illustrations, charts, tables; logos in the header or footer band
    and small icons left out). With visual_first, a slide with BARE_TEXT 字 or more and none of them gets a
    hint: a grid of labels and short phrases never trips the paragraph hints, so this is how it shows up.
    Information drawn with plain shapes (dots, bars, arrows) has no picture either: look at the image."""
    summary["visuals"] = visuals
    if visual_first and visuals == 0 and summary["units"] >= BARE_TEXT:
        summary["findings"].append({
            "kind": "no_picture_slide", "severity": "hint", "text": "(whole slide)", "rect": None,
            "units": summary["units"],
            "detail": f"About {summary['units']} 字 and no picture, illustration, chart or table detected. Look at the "
                      "slide: draw what can be seen, counted or followed (things, places, numbers, steps, flows) with "
                      "an illustration or photo, icons with short labels, dots, bars, arrows. Shapes that already draw "
                      "it count: say so. Do not add a decorative picture just to clear the hint."})
    return summary


def print_text(text):
    print("text (字 per slide, body only; in sentences): "
          + "  ".join(f"{s['slide']:02d}:{s['units']}({s['sentence_units']})" for s in text["slides"]))
    bare = [s["slide"] for s in text["slides"] if s.get("visuals") == 0 and s["units"] >= BARE_TEXT]
    if bare:
        print(f"no picture, illustration, chart or table on {len(bare)} of {len(text['slides'])} slides with "
              f"{BARE_TEXT}+ 字: " + ", ".join(f"{n:02d}" for n in bare))
    if text["visual_first"]:
        hints = [(s["slide"], f) for s in text["slides"] for f in s["findings"]]
        print(f"visual-first hints: {len(hints)} (handle each one: draw the information, then cut repeats or move "
              "explanations to notes; never drop items; not a pass/fail check)")
        for n, f in hints:
            print(f"  slide {n:02d} hint    {f['kind']:<18} “{f['text']}” {f['detail']}")


def draw_layout(shot, findings, out_path, scale):
    """Annotate a slide screenshot with numbered boxes: red = error, orange = warning."""
    Image = need("PIL.Image", "python -m pip install pillow")
    from PIL import ImageDraw
    im = Image.open(shot).convert("RGB")
    draw = ImageDraw.Draw(im)
    font = _font(max(14, int(16 * scale)))
    numbered = sorted(enumerate(findings, 1), key=lambda p: p[1]["severity"] == "error")  # errors drawn on top
    for n, f in numbered:
        r = f["rect"]
        color = (220, 30, 30) if f["severity"] == "error" else (240, 140, 0)
        box = [r["x"] * scale, r["y"] * scale, (r["x"] + r["w"]) * scale, (r["y"] + r["h"]) * scale]
        draw.rectangle(box, outline=color, width=max(2, int(3 * scale)))
        draw.text((box[0] + 2, max(0, box[1] - 20 * scale)), str(n), fill=color, font=font)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    im.save(out_path)


def check(name, status, detail="", required=True, **evidence):
    return {"name": name, "status": status, "required": required, "detail": detail, **evidence}


def checks_status(checks):
    required = [c["status"] for c in checks if c["required"]]
    for status in ("failed", "unsupported", "not_run"):
        if status in required:
            return status
    return "passed" if required and all(s == "passed" for s in required) else "not_run"


def write_report(report, out, strict=False):
    report["automatic_status"] = checks_status(report["checks"])
    report["manual_review"] = {
        "status": "not_run", "required_for_delivery": True,
        "detail": "View every final page; inspect content, visual quality and task-specific edits. "
                  "Automatic checks do not prove target-office interaction or design quality."
    }
    (out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if strict and report["automatic_status"] != "passed":
        raise SystemExit(1)


def launch_browser(p, choice):
    order = {"auto": [None, "chrome", "msedge"], "chromium": [None], "chrome": ["chrome"], "msedge": ["msedge"]}[choice]
    errors = []
    for channel in order:
        try:
            return p.chromium.launch(channel=channel) if channel else p.chromium.launch(), channel or "chromium"
        except Exception as e:  # browser not installed for this channel
            errors.append(f"{channel or 'chromium'}: {str(e).splitlines()[0]}")
    sys.exit("no usable browser:\n  " + "\n  ".join(errors) +
             "\ninstall one with: python -m playwright install chromium")


def cmd_html(args):
    sync_api = need("playwright.sync_api", "python -m pip install playwright && python -m playwright install chromium")
    src = Path(args.input).resolve()
    out = Path(args.out)
    slides_dir = out / "slides"
    if slides_dir.exists():
        shutil.rmtree(slides_dir)
    slides_dir.mkdir(parents=True)
    report = {"file": str(src), "console": [], "page_errors": [], "blocked_requests": [], "failed_requests": []}

    with sync_api.sync_playwright() as p:
        browser, used = launch_browser(p, args.browser)
        report["browser"] = used
        ctx = browser.new_context(viewport={"width": 1280, "height": 720}, device_scale_factor=args.scale)

        def route(r):
            url = r.request.url
            if url.startswith(("file:", "data:", "blob:", "about:")):
                r.continue_()
            else:
                report["blocked_requests"].append(url)
                r.abort()
        ctx.route("**/*", route)
        page = ctx.new_page()
        page.on("console", lambda m: m.type in ("error", "warning") and report["console"].append(f"{m.type}: {m.text}"))
        page.on("pageerror", lambda e: report["page_errors"].append(str(e)))
        page.on("requestfailed", lambda r: report["failed_requests"].append(f"{r.url} {r.failure}"))
        page.goto(src.as_uri(), wait_until="load")
        report["assets"] = page.evaluate(WAIT_ASSETS_JS)
        report["fonts_status"] = report["assets"]["fonts_status"]

        info = page.evaluate("""() => {
          const deck = document.getElementById('ssr-deck');
          return { width: Number(deck?.dataset.width || 1280), height: Number(deck?.dataset.height || 720),
                   slides: document.querySelectorAll('#ssr-deck > .ssr-slide').length,
                   api: typeof window.SSRPlayer === 'object' };
        }""")
        report.update(info)
        W, H, n = info["width"], info["height"], info["slides"]
        if not n:
            sys.exit("no #ssr-deck > .ssr-slide found")

        # Wide and small viewports with the player UI visible.
        page.set_viewport_size({"width": 1440, "height": 900})
        page.mouse.move(700, 450)
        page.wait_for_timeout(300)
        page.screenshot(path=str(out / "ui-wide.png"))
        page.set_viewport_size({"width": 800, "height": 600})
        page.mouse.move(400, 300)
        page.wait_for_timeout(300)
        page.screenshot(path=str(out / "ui-small.png"))

        # Keyboard navigation.
        if info["api"]:
            page.set_viewport_size({"width": W, "height": H})
            page.mouse.click(5, 5)
            checks = {}
            page.keyboard.press("End")
            checks["End -> last"] = page.evaluate("SSRPlayer.index") == n - 1
            page.keyboard.press("Home")
            checks["Home -> first"] = page.evaluate("SSRPlayer.index") == 0
            total = page.evaluate("SSRPlayer.totalSteps") + n + 2
            for _ in range(total):
                page.keyboard.press("ArrowRight")
            checks["ArrowRight x all -> last"] = page.evaluate("SSRPlayer.index") == n - 1
            for _ in range(total):
                page.keyboard.press("ArrowLeft")
            checks["ArrowLeft x all -> first"] = page.evaluate("SSRPlayer.index") == 0
            page.evaluate("location.hash = '#2'")
            page.wait_for_timeout(100)
            checks["hash #2 -> slide 2"] = n < 2 or page.evaluate("SSRPlayer.index") == 1
            report["keyboard"] = checks

        # Per-slide captures, fully revealed, without UI or animation.
        page.set_viewport_size({"width": W, "height": H})
        page.add_style_tag(content=CAPTURE_CSS)
        report["layout"] = []
        report["text"] = {"visual_first": args.visual_first, "slides": []}
        frames = []
        for i in range(n):
            if info["api"]:
                page.evaluate("i => SSRPlayer.go(i, { reveal: 'all' })", i)
            else:
                page.keyboard.press("Home")
                for _ in range(i):
                    page.keyboard.press("ArrowRight")
            page.wait_for_timeout(60)
            page.locator("#ssr-deck").screenshot(path=str(slides_dir / f"slide-{i + 1:02d}.png"))
            lint = page.evaluate(LAYOUT_JS, i)
            found = small_text([item[:4] for item in lint["sized"]], lint["height"])
            if found:
                count, total, sizes = found
                rects = [item[4] for item in lint["sized"] if is_small(item[1])
                         and lint["height"] * 0.115 <= item[2] and item[3] <= lint["height"] * 0.92]
                x0, y0 = min(r["x"] for r in rects), min(r["y"] for r in rects)
                x1, y1 = max(r["x"] + r["w"] for r in rects), max(r["y"] + r["h"] for r in rects)
                lint["findings"].append({"kind": "small_text", "severity": "warning", "text": "(body text)",
                                         "rect": {"x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0},
                                         "detail": small_text_detail(count, total, sizes, "px")})
            report["layout"].append({"slide": i + 1, "findings": lint["findings"]})
            frames.append((i + 1, page.evaluate(FRAME_JS, i)))
            body = page.evaluate(TEXT_JS, i)
            leaves = [{"text": l["text"], "rect": l["rect"], "groups": [l["group"]] if l["group"] >= 0 else []}
                      for l in body["leaves"]]
            report["text"]["slides"].append(bare_slide({"slide": i + 1, **text_density(leaves, body["groups"],
                                                                                       args.visual_first)},
                                                       body["visuals"], args.visual_first))

        if not args.no_pdf:
            pdf_path = out / "print.pdf"
            page.pdf(path=str(pdf_path), prefer_css_page_size=True, print_background=True)
            pm = pymupdf_or_none()
            if pm:
                doc = pm.open(pdf_path)
                r = doc[0].rect
                report["print"] = {"pages": doc.page_count, "expected_pages": n,
                                   "page_pt": [round(r.width, 1), round(r.height, 1)],
                                   "expected_pt": [round(W * 0.75, 1), round(H * 0.75, 1)]}
                doc.close()
            else:
                report["print"] = {"pdf": str(pdf_path), "note": "install PyMuPDF to count pages"}
        browser.close()

    shots = sorted(slides_dir.glob("slide-*.png"))
    make_sheet(shots, out / "sheet.png")
    zoom_dir = out / "zoom"
    if zoom_dir.exists():
        shutil.rmtree(zoom_dir)
    if args.zoom:
        for shot in shots:
            zoom_quarters(shot, zoom_dir, shot.stem)
    layout_dir = out / "layout"
    if layout_dir.exists():
        shutil.rmtree(layout_dir)
    for item in report["layout"]:
        if item["findings"]:
            draw_layout(slides_dir / f"slide-{item['slide']:02d}.png", item["findings"],
                        layout_dir / f"slide-{item['slide']:02d}.png", args.scale)
    frame = deck_frame(frames, "px", tolerance=3, threshold=16)
    report["frames"] = [{"slide": s, **f} if f else {"slide": s} for s, f in frames]
    report["deck_layout"] = [frame] if frame else []
    errors = sum(f["severity"] == "error" for item in report["layout"] for f in item["findings"])
    warnings = sum(f["severity"] == "warning" for item in report["layout"] for f in item["findings"]) + len(report["deck_layout"])
    problems = report["console"] + report["page_errors"] + report["failed_requests"]
    print(f"browser: {report['browser']}  slides: {n}  canvas: {W}x{H}  fonts: {report['fonts_status']}")
    print(f"slides -> {slides_dir}")
    print(f"sheet  -> {out / 'sheet.png'} (+ -gray, -blur)   ui -> ui-wide.png, ui-small.png")
    if "keyboard" in report:
        print("keyboard: " + ", ".join(f"{k}: {'ok' if v else 'FAIL'}" for k, v in report["keyboard"].items()))
    if "print" in report:
        print(f"print: {report['print']}")
    print(f"blocked external requests: {len(report['blocked_requests'])}  console/page errors: {len(problems)}")
    for line in (report["blocked_requests"] + problems)[:20]:
        print("  " + line)
    print(f"layout: {errors} errors, {warnings} warnings" + (f" -> annotated pages in {layout_dir}" if errors or warnings else ""))
    for item in report["layout"]:
        for f in item["findings"]:
            print(f"  slide {item['slide']:02d} {f['severity']:<7} {f['kind']:<18} “{f['text']}” {f['detail']}")
    for f in report["deck_layout"]:
        print(f"  deck     {f['severity']:<7} {f['kind']:<18} {f['detail']}")
    if args.zoom:
        print(f"zoom   -> {zoom_dir} (four enlarged quarters per slide: look again at text, breaks and edges)")
    print_text(report["text"])

    report["checks"] = [
        check("offline_resources", "failed" if report["blocked_requests"] or report["failed_requests"] else "passed"),
        check("browser_errors", "failed" if report["console"] or report["page_errors"] else "passed"),
        check("image_decode", "passed" if all(i["ok"] for i in report["assets"]["images"]) else "failed",
              "HTML img and SVG image elements; inspect CSS backgrounds and crop quality visually too."),
        check("font_load", "passed" if report["fonts_status"] == "loaded" and
              all(f["status"] != "error" for f in report["assets"]["fonts"]) else "failed",
              "Loading only; glyph coverage and fallback still require font checks and visual inspection."),
        check("keyboard", "passed" if report.get("keyboard") and all(report["keyboard"].values())
              else "failed" if report.get("keyboard") else "unsupported",
              "SSRPlayer keyboard checks; use task-specific checks for a different player."),
        check("layout", "failed" if errors else "passed",
              f"{errors} errors (outside slide, clipped, overlapping text), {warnings} warnings (orphan line, "
              "small font, contrast, placeholder, off-centre content frame). Geometry only: composition, meaning "
              "and crops still need review.",
              errors=errors, warnings=warnings),
    ]
    pr = report.get("print", {})
    if args.no_pdf:
        report["checks"].append(check("print", "not_run", "Skipped with --no-pdf; strict mode requires this check."))
    elif "pages" not in pr:
        report["checks"].append(check("print", "unsupported", "PyMuPDF needed for page count and dimensions."))
    else:
        ok = pr["pages"] == pr["expected_pages"] and all(
            abs(a-b) <= 1 for a,b in zip(pr["page_pt"], pr["expected_pt"]))
        report["checks"].append(check("print", "passed" if ok else "failed", **pr))
    write_report(report, out, args.strict)


# ---------------------------------------------------------------- PPTX

def find_soffice(explicit):
    """Path to LibreOffice's soffice, or None if it is not available (rendering is then reported as not_run)."""
    if explicit:
        found = explicit if Path(explicit).is_file() else shutil.which(explicit)
        if not found:
            print(f"warning: LibreOffice not found at {explicit}", file=sys.stderr)
        return found
    for name in ("soffice.com", "soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.com",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice", "/opt/libreoffice/program/soffice",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return None


PLACEHOLDER = re.compile(r"lorem ipsum|\bTODO\b|\bTBD\b|X{3,}|占位|待补充|待填写|此处填写|示例文字", re.I)
SYMBOLS = re.compile(r"^[\W_]+$")
DISPLAY_PT = 54  # text at least this large is display type (huge numbers, cover words) that may bleed on purpose


def is_small(px):
    return px < SMALL_PX - 0.5  # renderers draw 12pt chart text at about 11.8pt


def small_text(items, page_height):
    """items: [(text, size in canvas px, top, bottom)] on one slide, canvas px. Body text between the header
    band (top 11.5%) and the footer band (bottom 8%), below headline size, counts. Returns (small, total,
    sizes) in 字 when at least SMALL_MIN 字 and SMALL_SHARE of that body text is under SMALL_PX, else None."""
    body = [(text, size) for text, size, top, bottom in items
            if top >= page_height * 0.115 and bottom <= page_height * 0.92 and size < 36]
    total = sum(text_units(text) for text, _ in body)
    small = [(text, size) for text, size in body if is_small(size)]
    count = sum(text_units(text) for text, _ in small)
    if count >= SMALL_MIN and count >= SMALL_SHARE * total:
        return count, total, sorted({round(size, 1) for _, size in small})
    return None


def small_text_detail(count, total, sizes, unit):
    shown = ", ".join(f"{s:g}{unit}" for s in sizes[:4])
    return (f"About {count} of {total} 字 of body text is under 16px on a 1280-wide canvas (12pt on a 13.33-inch "
            f"slide): {shown}. Look at what was flagged. Body text and chart labels below the typography.md floor "
            "(talks 20px, reading 18px, labels 16px): enlarge, split the slide or draw it, never shrink to fit. "
            "Sources, page numbers, footnotes and photo credits may stay small (12px or more); note where and why.")


def thin_segments(drawings, max_width=6.0):
    """Straight strokes no wider than max_width, as (x0, y0, x1, y1) with the drawing's paint order.
    Hairline-thin filled rectangles count too: some renderers draw rules that way."""
    for d in drawings:
        if d.get("color") is not None and d.get("type") in ("s", "fs") and (d.get("width") or 0) <= max_width:
            for item in d["items"]:
                if item[0] == "l":
                    yield (item[1].x, item[1].y, item[2].x, item[2].y), d["seqno"]
                elif item[0] == "re":
                    r = item[1]
                    for edge in ((r.x0, r.y0, r.x1, r.y0), (r.x1, r.y0, r.x1, r.y1),
                                 (r.x0, r.y1, r.x1, r.y1), (r.x0, r.y0, r.x0, r.y1)):
                        yield edge, d["seqno"]
        if d.get("fill") is not None and d.get("type") in ("f", "fs"):
            r = d["rect"]
            if r.height <= 2.5 and r.width > 8:
                yield (r.x0, (r.y0 + r.y1) / 2, r.x1, (r.y0 + r.y1) / 2), d["seqno"]
            elif r.width <= 2.5 and r.height > 8:
                yield ((r.x0 + r.x1) / 2, r.y0, (r.x0 + r.x1) / 2, r.y1), d["seqno"]


def segment_hits(seg, box):
    """Whether the segment passes through the box (Liang-Barsky clipping)."""
    px, py, qx, qy = seg
    dx, dy = qx - px, qy - py
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, px - box[0]), (dx, box[2] - px), (-dy, py - box[1]), (dy, box[3] - py)):
        if p == 0:
            if q < 0:
                return False
        elif p < 0:
            t0 = max(t0, q / p)
        else:
            t1 = min(t1, q / p)
    return t0 < t1


def glyph_core(bbox, size):
    """The middle of a glyph cell: line-spacing slack, side bearings and the under/over-line zones removed."""
    x0, y0, x1, y1 = bbox
    inset = max(0.0, (y1 - y0 - size) / 2)
    y0, y1 = y0 + inset, y1 - inset
    dx, dy = max(0.5, (x1 - x0) * 0.15), (y1 - y0) * 0.25
    return x0 + dx, y0 + dy, x1 - dx, y1 - dy


def lines_through_text(page, flags):
    """Thin drawn lines that run through the middle of glyphs, as (text, span bbox). A filled shape painted
    after the line over the glyph hides it; a horizontal stroke spanning exactly the run is a strikethrough."""
    drawings = page.get_drawings()
    segments = list(thin_segments(drawings))
    if not segments:
        return []
    fills = [(d["seqno"], d["rect"]) for d in drawings if d.get("fill") is not None]

    def hidden(core, seq):
        return any(s > seq and r.x0 <= core[0] and r.y0 <= core[1] and r.x1 >= core[2] and r.y1 >= core[3]
                   for s, r in fills)

    def strikethrough(seg, span_box):
        return (abs(seg[1] - seg[3]) < 0.1 and abs(min(seg[0], seg[2]) - span_box[0]) < 1.5
                and abs(max(seg[0], seg[2]) - span_box[2]) < 1.5)

    hits = []
    for block in page.get_text("rawdict", flags=flags)["blocks"]:
        if block.get("type") != 0:
            continue
        for span in (s for line in block["lines"] for s in line["spans"]):
            text = "".join(c["c"] for c in span["chars"]).strip()
            if not text or SYMBOLS.match(text):
                continue
            sx0, sy0, sx1, sy1 = span["bbox"]
            near = [(seg, seq) for seg, seq in segments
                    if min(seg[0], seg[2]) <= sx1 and max(seg[0], seg[2]) >= sx0
                    and min(seg[1], seg[3]) <= sy1 and max(seg[1], seg[3]) >= sy0
                    and not strikethrough(seg, span["bbox"])]
            crossed = False
            for ch in span["chars"]:
                if not near or crossed:
                    break
                if not ch["c"].strip() or SYMBOLS.match(ch["c"]):
                    continue
                core = glyph_core(ch["bbox"], span["size"])
                crossed = any(segment_hits(seg, core) and not hidden(core, seq) for seg, seq in near)
            if crossed:
                hits.append((text, span["bbox"]))
    return hits


def lone_character(text):
    """A line holding one Chinese/Japanese/Korean character, punctuation aside: "间；", "性。", "在".
    A lone digit or letter is usually a label of its own (a big number, an axis tick), so it does not count."""
    core = [c for c in text if c.isalnum()]
    return len(core) == 1 and bool(CJK_CHAR.match(core[0]))


def boxes(drawings, width, height):
    """Filled or outlined shapes big enough to hold text (cards, panels, bands), as rects in points."""
    out = []
    for d in drawings:
        r = d["rect"]
        if r.width < 24 or r.height < 14 or r.width * r.height > 0.5 * width * height:
            continue
        filled = d.get("fill") is not None and d.get("type") in ("f", "fs")
        stroked = d.get("color") is not None and d.get("type") in ("s", "fs")
        if filled or stroked:
            out.append(r)
    return out


def text_groups(page, flags, containers=()):
    """Paragraph-like groups of rendered text. PDF text extraction joins everything on one baseline, even
    labels in separate boxes, so lines are first cut where the gap between characters exceeds 1.5 em; the
    pieces are then stacked into one group when they share a left, right or centre edge, have about the
    same size, sit less than 0.9 em apart and keep the group's line pitch (stacked labels in separate boxes
    usually do not). A line that sits in a box the line above is not in (the next table cell or card)
    starts a new group. Each group: {"lines": [(text, bbox)], "bbox", "size"}."""

    def holds(r, b):
        return r.x0 - 1 <= b[0] and r.y0 - 1 <= b[1] and r.x1 + 1 >= b[2] and r.y1 + 1 >= b[3]
    pieces = []
    for block in page.get_text("rawdict", flags=flags)["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            if line.get("dir", (1, 0))[1] != 0:
                continue  # rotated or vertical text
            chars = sorted(((c, s["size"]) for s in line["spans"] for c in s["chars"]), key=lambda t: t[0]["bbox"][0])
            current, last = [], None
            for c, size in chars:
                if c["c"].strip() and last is not None and c["bbox"][0] - last["bbox"][2] > 1.5 * size:
                    pieces.append((current, size))
                    current = []
                current.append(c)
                if c["c"].strip():
                    last = c
            if current:
                pieces.append((current, max(s["size"] for s in line["spans"])))
    segments = []
    for chars, size in pieces:
        ink = [c for c in chars if c["c"].strip()]
        if not ink:
            continue
        text = "".join(c["c"] for c in chars).strip()
        bbox = (min(c["bbox"][0] for c in ink), min(c["bbox"][1] for c in ink),
                max(c["bbox"][2] for c in ink), max(c["bbox"][3] for c in ink))
        segments.append({"text": text, "bbox": bbox, "size": size})
    segments.sort(key=lambda s: (s["bbox"][1], s["bbox"][0]))
    groups = []
    for seg in segments:
        x0, y0, x1, y1 = seg["bbox"]
        size = seg["size"]
        home = None
        for g in groups:
            lx0, ly0, lx1, ly1 = g["lines"][-1][1]
            tol = max(2.0, 0.3 * size)
            aligned = abs(lx0 - x0) <= tol or abs(lx1 - x1) <= tol or abs((lx0 + lx1) - (x0 + x1)) / 2 <= tol
            pitch = y0 - ly0
            steady = len(g["lines"]) < 2 or abs(pitch - (ly0 - g["lines"][-2][1][1])) <= 0.25 * pitch
            if (aligned and steady and abs(g["size"] - size) <= 0.15 * size and pitch > 0.5 * size
                    and y0 - ly1 <= 0.9 * size
                    and not any(holds(r, seg["bbox"]) and not holds(r, (lx0, ly0, lx1, ly1)) for r in containers)):
                home = g
                break
        if home is None:
            groups.append({"lines": [(seg["text"], seg["bbox"])], "size": size, "bbox": list(seg["bbox"])})
        else:
            home["lines"].append((seg["text"], seg["bbox"]))
            b = home["bbox"]
            home["bbox"] = [min(b[0], x0), min(b[1], y0), max(b[2], x1), max(b[3], y1)]
    return groups


def page_frame(page, blocks, drawings):
    """Left and right margins of the content between the header band (top 11.5%) and the footer band
    (bottom 8%), leaving out full-width bands; None when nothing is there."""
    width, height = page.rect.width, page.rect.height
    rects = [d["rect"] for d in drawings]
    rects += [s["bbox"] for b in blocks for line in b["lines"] for s in line["spans"] if s["text"].strip()]
    rects += [info["bbox"] for info in page.get_image_info()]
    xs = [(r[0], r[2]) for r in rects
          if r[2] - r[0] < width * 0.95 and r[1] >= height * 0.115 and r[3] <= height * 0.92
          and (r[2] - r[0] >= 0.5 or r[3] - r[1] >= 0.5)]
    if not xs:
        return None
    left, right = min(x0 for x0, _ in xs), max(x1 for _, x1 in xs)
    return {"left": round(left, 1), "right": round(width - right, 1), "span": round(right - left, 1), "width": width}


def pdf_layout(page, pm, frame=None):
    """Text geometry on a rendered PPTX page, in points: text off the page, text over other text, a line of
    about one character, text running out of its card or panel, tiny text, placeholders and drawn lines
    through text. Backgrounds are not known, so no contrast. frame: a dict that receives the page's
    content margins (see page_frame)."""
    flags = getattr(pm, "TEXTFLAGS_DICT", 0) & ~getattr(pm, "TEXT_MEDIABOX_CLIP", 64)  # keep text past the edge
    width, height = page.rect.width, page.rect.height
    findings, spans = [], []
    drawings = page.get_drawings()
    containers = boxes(drawings, width, height)

    def add(kind, severity, text, bbox, detail):
        x0, y0, x1, y1 = bbox
        findings.append({"kind": kind, "severity": severity, "text": text.strip()[:40],
                         "rect": {"x": round(x0), "y": round(y0), "w": round(x1 - x0), "h": round(y1 - y0)},
                         "detail": detail})

    blocks = [b for b in page.get_text("dict", flags=flags)["blocks"] if b.get("type") == 0]
    if frame is not None:
        frame.update(page_frame(page, blocks, drawings) or {})
    for bi, block in enumerate(blocks):
        lines = []
        for li, line in enumerate(block["lines"]):
            text = "".join(s["text"] for s in line["spans"])
            if text.strip():
                lines.append((text.strip(), line["bbox"]))
            for s in line["spans"]:
                if s["text"].strip():
                    spans.append((bi, li, s))
        if not lines:
            continue
        block_text = " ".join(t for t, _ in lines)
        if PLACEHOLDER.search(block_text):
            add("placeholder", "warning", block_text, block["bbox"], "Looks like placeholder text.")
        tiny = min((s["size"] for line in block["lines"] for s in line["spans"] if s["text"].strip()), default=99)
        if tiny < 9:
            add("font_too_small", "warning", block_text, block["bbox"], f"{tiny:.1f}pt (< 9pt, about 12px).")

    for g in text_groups(page, flags, containers):
        gx0, gy0, gx1, gy1 = g["bbox"]
        if gx0 < -1 or gy0 < -1 or gx1 > width + 1 or gy1 > height + 1:
            continue  # reported below as an error
        group_text = " ".join(t for t, _ in g["lines"])
        # A line of one character inside a paragraph: a wrapped tail or the word before a forced break.
        # Vertical East Asian text comes out as a stack of one-character lines; that is not an orphan.
        vertical = all(sum(c.isalnum() for c in t) <= 1 for t, _ in g["lines"])
        for text, bbox in ([] if vertical else g["lines"][1:]):
            if lone_character(text):
                add("orphan_line", "warning", group_text, bbox,
                    "A line holds one character; rewrite, widen the box or break the line deliberately.")
                break
        # Text that starts inside a card or panel and runs past its edge: no box around the first line holds
        # the whole paragraph (a heading bar inside a larger card is fine).
        first = g["lines"][0][1]
        inside = [r for r in containers if r.x0 - 1 <= first[0] and r.y0 - 1 <= first[1]
                  and r.x1 + 1 >= first[2] and r.y1 + 1 >= first[3]]
        if inside and not any(gy1 <= r.y1 + 2 and gx1 <= r.x1 + 2 and gx0 >= r.x0 - 2 for r in inside):
            add("text_outside_box", "warning", group_text, g["bbox"],
                "Text starts inside a card or panel and runs past its edge; shorten the text or enlarge the box.")

    def em_box(s):
        x0, y0, x1, y1 = s["bbox"]
        inset = max(0.0, (y1 - y0 - s["size"]) / 2)
        return x0, y0 + inset, x1, y1 - inset

    for bi, li, s in spans:
        x0, y0, x1, y1 = em_box(s)  # the glyphs' em box: line-spacing slack above a huge glyph is not text
        if x0 < -1 or y0 < -1 or x1 > width + 1 or y1 > height + 1:
            if s["size"] >= DISPLAY_PT:
                add("text_outside_slide", "warning", s["text"], s["bbox"],
                    "Display text runs past the slide edge. Fine for an intended bleed (a huge number or word "
                    "cropped by the edge); otherwise move it in.")
            else:
                add("text_outside_slide", "error", s["text"], s["bbox"], "Text runs past the slide edge.")

    # Compare every pair of spans: text extraction may merge overlapping boxes into one block or line.
    # Neighbouring spans and consecutive lines only touch; em-box height ignores line-spacing slack.
    for i, (bi, li, a) in enumerate(spans):
        if SYMBOLS.match(a["text"].strip()):
            continue
        A = em_box(a)
        for bj, lj, b in spans[i + 1:]:
            if SYMBOLS.match(b["text"].strip()):
                continue
            B = em_box(b)
            if min(A[2], B[2]) - max(A[0], B[0]) > 2 and min(A[3], B[3]) - max(A[1], B[1]) > 2:
                if a["text"].strip() == b["text"].strip():
                    add("text_overlap", "warning", a["text"], a["bbox"],
                        "The same text is drawn twice, offset (a shadow or misprint effect). Fine when intended; "
                        "otherwise delete the copy.")
                else:
                    add("text_overlap", "error", a["text"], a["bbox"], f"Overlaps “{b['text'].strip()[:40]}”.")
                break

    for text, bbox in lines_through_text(page, flags):
        add("line_through_text", "warning", text, bbox,
            "A drawn line runs through this text; move the line or the label, or give the label a background. "
            "An intentional strike or decoration can stay.")

    # Small body text, judged on a 1280-wide canvas so 4:3 and 16:9 slides of any size compare alike.
    unit = 1280 / width
    found = small_text([(s["text"], s["size"] * unit, s["bbox"][1] * unit, s["bbox"][3] * unit) for _, _, s in spans],
                       height * unit)
    if found:
        small = [s for _, _, s in spans if is_small(s["size"] * unit)
                 and height * 0.115 <= s["bbox"][1] and s["bbox"][3] <= height * 0.92]
        box = (min(s["bbox"][0] for s in small), min(s["bbox"][1] for s in small),
               max(s["bbox"][2] for s in small), max(s["bbox"][3] for s in small))
        count, total, sizes = found
        add("small_text", "warning", small[0]["text"], box,
            small_text_detail(count, total, sorted({round(z / unit, 1) for z in sizes}), "pt"))

    # What the rendering shows: text that leaves no ink (covered by a shape or picture drawn later, or set in
    # the background colour) and text with too little contrast against what is painted behind it.
    # These are clues: a clean result does not prove that every word can be seen and read (see visual-review).
    pixmap = page.get_pixmap(matrix=pm.Matrix(INK_ZOOM, INK_ZOOM), alpha=False)
    picture = Image_from_pixmap(pixmap)
    hidden = hidden_text(page, flags, picture, INK_ZOOM)
    for text, bbox, count, total in hidden:
        add("text_hidden", "error", text, bbox,
            f"{count} of {total} characters leave no ink in the rendering: a shape or picture drawn later covers "
            "them, or they are set in the background colour. Bring the text forward or move the shape.")
    for text, bbox in text_under_picture(page, pm, flags, picture, INK_ZOOM, [b for _, b, _, _ in hidden]):
        add("text_under_picture", "warning", text, bbox,
            "A picture is drawn on top of this text. Look at the page: if the picture hides the words (a texture or "
            "pattern can look like ink to the hidden-text check), bring the text forward or move the picture.")
    for (fg, bg, large), (ratio, chars, text, bbox) in low_contrast(page, picture, INK_ZOOM).items():
        add("low_contrast", "warning", text, bbox,
            f"Contrast {ratio:.2f}:1 for #{fg} on #{bg} (needs {3 if large else 4.5}:1 for "
            f"{'large' if large else 'body'} text); {chars} 字 fall below it. Darken the text or lighten the "
            "background. Measured per character from the rendering; over photos and gradients the estimate is "
            "lenient, so look at such text yourself.")
    return findings


INK_ZOOM = 2.0  # rendering scale for the ink and contrast checks (2 px per point)


def Image_from_pixmap(pixmap):
    Image = need("PIL.Image", "python -m pip install pillow")
    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def hidden_text(page, flags, picture, zoom):
    """Spans whose glyph cores show (almost) no variation in the rendering: (text, bbox, hidden, total).
    Reported when at least 3 characters and half of the span's letters and digits leave no ink."""
    gray = picture.convert("L")
    out = []
    for block in page.get_text("rawdict", flags=flags)["blocks"]:
        if block.get("type") != 0:
            continue
        for span in (s for line in block["lines"] for s in line["spans"]):
            chars = [c for c in span["chars"] if c["c"].isalnum()]
            if len(chars) < 2:
                continue
            hidden = 0
            for c in chars:
                x0, y0, x1, y1 = glyph_core(c["bbox"], span["size"])
                box = (int(x0 * zoom), int(y0 * zoom), max(int(x0 * zoom) + 1, int(x1 * zoom)),
                       max(int(y0 * zoom) + 1, int(y1 * zoom)))
                if box[0] < 0 or box[1] < 0 or box[2] > gray.width or box[3] > gray.height:
                    continue  # off the page: reported as text outside the slide
                lo, hi = gray.crop(box).getextrema()
                hidden += hi - lo < 24
            if hidden >= 3 and hidden >= 0.5 * len(chars):
                out.append(("".join(c["c"] for c in span["chars"]).strip(), span["bbox"], hidden, len(chars)))
    return out


def relative_luminance(rgb):
    def channel(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = rgb
    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def contrast_ratio(a, b):
    hi, lo = sorted((relative_luminance(a), relative_luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def text_under_picture(page, pm, flags, picture, zoom, skip=()):
    """Text that a picture drawn later lies on top of: (text, bbox). The paint order gives the candidates (a
    picture painted after the text over most of its box); full-page pictures (grain, texture overlays) are
    left out. The paint log holds a cropped picture's full, uncropped box, so each candidate is confirmed by
    rendering a copy of the page without that picture: only text whose pixels change is reported. Text already
    reported as hidden is skipped."""
    log = page.get_bboxlog()
    area = page.rect.width * page.rect.height
    pictures = [(i, pm.Rect(r)) for i, (kind, r) in enumerate(log)
                if kind in ("fill-image", "fill-imgmask") and pm.Rect(r).get_area() < 0.85 * area]
    covered = []  # (text rect, picture rects drawn over it)
    for i, (kind, r) in enumerate(log):
        rect = pm.Rect(r)
        if kind not in ("fill-text", "stroke-text") or rect.is_empty:
            continue
        over = [p for j, p in pictures if j > i and (rect & p).get_area() >= 0.6 * rect.get_area()]
        if over:
            covered.append((rect, over))
    candidates = []
    for block in page.get_text("dict", flags=flags)["blocks"]:
        for span in (s for line in block.get("lines", []) for s in line["spans"]):
            box = pm.Rect(span["bbox"])
            text = span["text"].strip()
            if not text or SYMBOLS.match(text) or box.is_empty or any(pm.Rect(b).intersects(box) for b in skip):
                continue
            over = [p for c, ps in covered if (box & c).get_area() >= 0.5 * box.get_area() for p in ps]
            if over:
                candidates.append((text, span["bbox"], over))
    if not candidates:
        return []
    xrefs = {}
    for info in page.get_image_info(xrefs=True):
        rect = pm.Rect(info["bbox"])
        for _, _, over in candidates:
            for p in over:
                if info.get("xref") and abs(rect.x0 - p.x0) + abs(rect.y0 - p.y0) + abs(rect.x1 - p.x1) + abs(rect.y1 - p.y1) < 2:
                    xrefs[info["xref"]] = rect
    if not xrefs:
        return [(text, bbox) for text, bbox, _ in candidates]  # cannot confirm: report for a look
    copy = pm.open("pdf", page.parent.tobytes())
    try:
        bare = copy[page.number]
        for xref in xrefs:
            bare.delete_image(xref)
        without = Image_from_pixmap(bare.get_pixmap(matrix=pm.Matrix(zoom, zoom), alpha=False)).convert("L")
    finally:
        copy.close()
    shown = picture.convert("L")
    out = []
    for text, bbox, _ in candidates:
        x0, y0, x1, y1 = (int(v * zoom) for v in bbox)
        box = (max(0, x0), max(0, y0), min(shown.width, max(x0 + 1, x1)), min(shown.height, max(y0 + 1, y1)))
        a, b = list(shown.crop(box).getdata()), list(without.crop(box).getdata())
        if a and sum(abs(p - q) for p, q in zip(a, b)) / len(a) > 6:
            out.append((text, bbox))
    return out


def low_contrast(page, picture, zoom):
    """Text colour against what is painted behind each character, read from the rendering: inside each
    character's box the darkest or lightest pixel, whichever is farther from the text colour, stands for the
    background (on a flat fill that is the fill itself; over a photo or gradient it is the most favourable
    pixel, so the estimate is lenient). A span is reported when at least 3 characters and a quarter of them
    fall below 4.5:1 (3:1 for large text), so a sentence that runs from a light onto a dark panel is caught.
    Only body text between the header and footer bands. Returns {(fg, bg, large): (ratio, characters, first
    text, bbox)}, with the worst ratio and its background per colour pair."""
    width, height = page.rect.width, page.rect.height
    unit = 1280 / width
    gray = picture.convert("L")
    luma = lambda rgb: round(rgb[0] * 0.299 + rgb[1] * 0.587 + rgb[2] * 0.114)
    found = {}
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                text = "".join(c["c"] for c in span["chars"]).strip()
                x0, y0, x1, y1 = span["bbox"]
                if len(text) < 2 or SYMBOLS.match(text) or y0 < height * 0.115 or y1 > height * 0.92:
                    continue
                c = span["color"]
                fg = ((c >> 16) & 255, (c >> 8) & 255, c & 255)
                size = span["size"] * unit
                large = size >= 24 or (size >= 18.66 and bool(span["flags"] & 16))
                need_ratio = 3 if large else 4.5
                measured, failing, worst = 0, 0, None
                for ch in span["chars"]:
                    if not ch["c"].strip() or SYMBOLS.match(ch["c"]):
                        continue
                    a, b, cx, cy = (int(v * zoom) for v in ch["bbox"])
                    box = (max(0, a), max(0, b), min(gray.width, max(a + 2, cx)), min(gray.height, max(b + 2, cy)))
                    if box[0] >= box[2] or box[1] >= box[3]:
                        continue
                    lo, hi = gray.crop(box).getextrema()
                    level = hi if abs(hi - luma(fg)) > abs(lo - luma(fg)) else lo
                    measured += 1
                    if contrast_ratio((level,) * 3, (luma(fg),) * 3) >= need_ratio + 1:
                        continue  # clearly fine on the grey estimate; skip the exact colour lookup
                    measured -= 1
                    data = list(gray.crop(box).getdata())
                    spot = data.index(level)
                    bw = box[2] - box[0]
                    bg = picture.getpixel((box[0] + spot % bw, box[1] + spot // bw))
                    ratio = contrast_ratio(fg, bg)
                    measured += 1
                    if ratio < need_ratio:
                        failing += 1
                        if worst is None or ratio < worst[0]:
                            worst = (ratio, bg)
                if worst and failing >= 3 and failing >= 0.25 * measured:
                    key = ("%02X%02X%02X" % fg, "%02X%02X%02X" % tuple(worst[1]), large)
                    old = found.get(key)
                    found[key] = (worst[0], failing, text, span["bbox"]) if old is None else \
                        (min(old[0], worst[0]), old[1] + failing, old[2], old[3])
    return found


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def pdf_to_pngs(pdf, out_dir, width, layout=None, zoom_dir=None):
    """One PNG per PDF page; with layout, also the layout findings and content margins of each page;
    with zoom_dir, four quarters of each page rendered at twice the width."""
    pm = pymupdf_or_none() or sys.exit("missing dependency PyMuPDF: python -m pip install pymupdf")
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = pm.open(pdf)
    paths = []
    try:
        for i, page in enumerate(doc):
            zoom = width / page.rect.width
            pix = page.get_pixmap(matrix=pm.Matrix(zoom, zoom))
            path = out_dir / f"slide-{i + 1:02d}.png"
            pix.save(path)
            paths.append(path)
            if layout is not None:
                frame = {}
                layout.append({"slide": i + 1, "scale": zoom, "findings": pdf_layout(page, pm, frame),
                               "frame": frame or None})
            if zoom_dir is not None:
                with tempfile.TemporaryDirectory() as tmp:
                    big = Path(tmp, "page.png")
                    page.get_pixmap(matrix=pm.Matrix(2 * zoom, 2 * zoom)).save(big)
                    zoom_quarters(big, zoom_dir, path.stem)
    finally:
        doc.close()
    return paths


def pptx_to_pngs(pptx, out_dir, soffice, width, layout=None, zoom_dir=None):
    pymupdf_or_none() or sys.exit("missing dependency PyMuPDF: python -m pip install pymupdf")
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "profile").resolve().as_uri()
        cmd = [soffice, f"-env:UserInstallation={profile}", "--headless", "--norestore",
               "--convert-to", "pdf", "--outdir", tmp, str(pptx)]
        subprocess.run(cmd, check=True, capture_output=True, timeout=300)
        pdf = Path(tmp, Path(pptx).stem + ".pdf")
        if not pdf.exists():
            raise RuntimeError(f"LibreOffice did not produce a PDF for {pptx}")
        return pdf_to_pngs(pdf, out_dir, width, layout, zoom_dir)


# Opens the deck read-only and windowless in WPS or PowerPoint through COM and saves a PDF
# (ppSaveAsPDF = 32), then closes only that presentation. The app is quit only when no WPS or PowerPoint
# process was running beforehand (WPS keeps all its documents in one process), so the user's open files are
# never touched. If the ProgID does not resolve in this registry view, the CLSID it names is tried.
EXPORT_PS1 = r"""param([string]$ProgId, [string]$Source, [string]$Target)
$ErrorActionPreference = 'Stop'
$running = @(Get-Process -Name wps, wpp, et, wpspdf, POWERPNT -ErrorAction SilentlyContinue).Count -gt 0
try { $app = New-Object -ComObject $ProgId }
catch {
  $clsid = (Get-ItemProperty "Registry::HKEY_CLASSES_ROOT\$ProgId\CLSID" -ErrorAction SilentlyContinue).'(default)'
  if (-not $clsid) { $clsid = (Get-ItemProperty "Registry::HKEY_CURRENT_USER\Software\Classes\$ProgId\CLSID" -ErrorAction SilentlyContinue).'(default)' }
  if (-not $clsid) { throw }
  $app = [Activator]::CreateInstance([Type]::GetTypeFromCLSID($clsid))
}
$pres = $null
try {
  $pres = $app.Presentations.Open($Source, -1, 0, 0)
  $pres.SaveAs($Target, 32)
} finally {
  if ($pres -ne $null) { $pres.Close() }
  if (-not $running) { try { $app.Quit() } catch {} }
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) | Out-Null
}
"""
EXPORT_APPS = {"wps": "KWPP.Application", "powerpoint": "PowerPoint.Application"}


def export_pdf(pptx, target, app):
    """Have WPS or PowerPoint (Windows, through COM) export the deck to target. Tries 64-bit and then
    32-bit PowerShell, since WPS often registers only a 32-bit COM server. Returns the PDF path."""
    if os.name != "nt":
        raise RuntimeError("--export needs Windows with WPS or PowerPoint installed; export the PDF yourself and use --pdf")
    windir = Path(os.environ.get("SystemRoot", r"C:\Windows"))
    shells = [p for p in (windir / "System32/WindowsPowerShell/v1.0/powershell.exe",
                          windir / "SysWOW64/WindowsPowerShell/v1.0/powershell.exe") if p.exists()]
    if not shells and shutil.which("powershell"):
        shells = [Path(shutil.which("powershell"))]
    if not shells:
        raise RuntimeError("Windows PowerShell not found")
    target = Path(target).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        target.unlink()
    errors = []
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp, "export.ps1")
        script.write_text(EXPORT_PS1, encoding="utf-8-sig")
        for shell in shells:
            proc = subprocess.run([str(shell), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                                   str(script), "-ProgId", EXPORT_APPS[app], "-Source", str(Path(pptx).resolve()),
                                   "-Target", str(target)], capture_output=True, text=True, timeout=300)
            if proc.returncode == 0 and target.exists():
                return target
            errors.append(f"{shell.parent.parent.parent.name}: {(proc.stderr or proc.stdout).strip()[:300]}")
    raise RuntimeError(f"{app} could not export the PDF ({'; '.join(errors)})")


NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def font_inventory(pptx):
    """Typefaces actually set on slides and in the theme (latin vs East Asian)."""
    used = {"latin": set(), "ea": set()}
    themes = {}
    with zipfile.ZipFile(pptx) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = ET.fromstring(z.read(name))
                for kind in ("latin", "ea"):
                    for el in root.iter(f"{{{NS['a']}}}{kind}"):
                        if el.get("typeface"):
                            used[kind].add(el.get("typeface"))
            elif name.startswith("ppt/theme/theme") and name.endswith(".xml"):
                root = ET.fromstring(z.read(name))
                theme = {}
                for role in ("majorFont", "minorFont"):
                    el = root.find(f".//a:{role}", NS)
                    if el is not None:
                        theme[role] = {k: (el.find(f"a:{k}", NS).get("typeface") if el.find(f"a:{k}", NS) is not None else None)
                                       for k in ("latin", "ea")}
                themes[name] = theme
    return {"slides_latin": sorted(used["latin"]), "slides_ea": sorted(used["ea"]), "themes": themes}


def inventory(pptx):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return None
    prs = Presentation(pptx)
    area = prs.slide_width * prs.slide_height
    slides = []

    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(s.shapes)
            else:
                yield s
    for idx, slide in enumerate(prs.slides, 1):
        c = {"slide": idx, "text_shapes": 0, "tables": 0, "charts": 0, "pictures": 0, "notes": False}
        big_picture = False
        for s in walk(slide.shapes):
            if s.has_text_frame and s.text_frame.text.strip():
                c["text_shapes"] += 1
            if getattr(s, "has_table", False) and s.has_table:
                c["tables"] += 1
            if getattr(s, "has_chart", False) and s.has_chart:
                c["charts"] += 1
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                c["pictures"] += 1
                if s.width and s.height and s.width * s.height >= 0.85 * area:
                    big_picture = True
        c["notes"] = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
        if big_picture and not (c["text_shapes"] or c["tables"] or c["charts"]):
            c["warning"] = "full-slide picture with no native text/table/chart"
        slides.append(c)
    return slides


def pptx_text(pptx, visual_first):
    """Body text per slide, read from the file (no rendering needed). Paragraphs are the leaves; a text box
    and a card shape underneath several text boxes are the enclosing blocks. Left out: titles, slide
    number/date/footer placeholders, boxes whose text is all headline size (>= 27pt), tables, charts and
    the footer band (bottom 12%). Rects are in points."""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
        from pptx.oxml.ns import qn
    except ImportError:
        return None
    prs = Presentation(pptx)
    W, H = prs.slide_width, prs.slide_height
    skip_ph = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE,
               PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER}

    def placed(shapes, place):
        """Yield (shape, (x, y, w, h)) in slide EMU, mapping group child coordinates."""
        for s in shapes:
            if s.left is None or s.top is None or s.width is None or s.height is None:
                continue
            box = place(s.left, s.top, s.width, s.height)
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                xfrm = s._element.grpSpPr.find(qn("a:xfrm"))
                off = xfrm.find(qn("a:chOff")) if xfrm is not None else None
                ext = xfrm.find(qn("a:chExt")) if xfrm is not None else None
                if off is None or ext is None or not int(ext.get("cx")) or not int(ext.get("cy")):
                    inner = place
                else:
                    ox, oy, cx, cy = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
                    gx, gy, gw, gh = box
                    inner = (lambda gx, gy, sx, sy, ox, oy: lambda x, y, w, h:
                             (gx + (x - ox) * sx, gy + (y - oy) * sy, w * sx, h * sy))(gx, gy, gw / cx, gh / cy, ox, oy)
                yield from placed(s.shapes, inner)
            else:
                yield s, box

    pt = lambda b: {"x": round(b[0] / 12700), "y": round(b[1] / 12700), "w": round(b[2] / 12700), "h": round(b[3] / 12700)}
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        leaves, groups, frames, cards = [], [], [], []
        visuals = 0  # pictures (not logos in the header or footer band), charts and tables
        for s, box in placed(slide.shapes, lambda x, y, w, h: (x, y, w, h)):
            is_ph = getattr(s, "is_placeholder", False)
            if is_ph and s.placeholder_format.type in skip_ph:
                continue
            area = box[2] * box[3]
            middle = box[1] + box[3] / 2
            if (getattr(s, "has_chart", False) and s.has_chart) or (getattr(s, "has_table", False) and s.has_table):
                visuals += 1
            elif (s.shape_type == MSO_SHAPE_TYPE.PICTURE or type(s).__name__ == "PlaceholderPicture") \
                    and area >= 0.004 * W * H and 0.115 * H <= middle <= 0.92 * H:
                visuals += 1
            if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 0.02 * W * H <= area <= 0.4 * W * H:
                cards.append(box)
            if not s.has_text_frame or not s.text_frame.text.strip() or box[1] >= 0.88 * H:
                continue
            sizes = [(r.font.size or p.font.size) for p in s.text_frame.paragraphs for r in p.runs if r.text.strip()]
            if sizes and all(z is not None and z.pt >= 27 for z in sizes):  # headline or big number
                continue
            frame = len(groups)
            groups.append(pt(box))
            frames.append((frame, box))
            for p in s.text_frame.paragraphs:
                for line in re.split(r"[\v\n]", p.text):
                    if line.strip():
                        leaves.append({"text": " ".join(line.split()), "rect": pt(box), "groups": [frame], "box": box})
        for card in cards:  # a card shape with two or more text boxes sitting on it
            inside = [f for f, b in frames if card[0] <= b[0] + b[2] / 2 <= card[0] + card[2]
                      and card[1] <= b[1] + b[3] / 2 <= card[1] + card[3]]
            if len(inside) >= 2:
                g = len(groups)
                groups.append(pt(card))
                for leaf in leaves:
                    if leaf["groups"][0] in inside:
                        leaf["groups"].append(g)
        for leaf in leaves:
            del leaf["box"]
        slides.append(bare_slide({"slide": idx, **text_density(leaves, groups, visual_first)}, visuals, visual_first))
    return {"visual_first": visual_first, "slides": slides}


def slide_paragraphs(slide):
    """Non-empty paragraph texts on a slide, including groups and table cells."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []

    def visit(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(s.shapes)
                continue
            frames = [s.text_frame] if s.has_text_frame else []
            if getattr(s, "has_table", False) and s.has_table:
                frames += [cell.text_frame for row in s.table.rows for cell in row.cells]
            for tf in frames:
                out.extend(" ".join(p.text.split()) for p in tf.paragraphs if p.text.strip())
    visit(slide.shapes)
    return out


def longest_common(a, b):
    """Length and text of the longest common substring of a and b."""
    best, end = 0, 0
    prev = [0] * (len(b) + 1)
    for i, ca in enumerate(a, 1):
        cur = [0] * (len(b) + 1)
        for j, cb in enumerate(b, 1):
            if ca == cb:
                cur[j] = prev[j - 1] + 1
                if cur[j] > best:
                    best, end = cur[j], i
        prev = cur
    return best, a[end - best:end]


def template_sample_text(pptx, template, min_len=5):
    """Text the deck still shares with the sample slides of the company template: department names, dates,
    sample titles left in place. Text that sits on most template slides (a fixed footer, a slogan) is part of
    the template and is left out, and so are shared runs under min_len characters. Returns
    [{"slide", "text", "template_text", "shared"}]."""
    from pptx import Presentation
    samples = [set(slide_paragraphs(s)) for s in Presentation(template).slides]
    seen = {}
    for paras in samples:
        for p in paras:
            seen[p] = seen.get(p, 0) + 1
    # Text on more than half of the template pages (and on at least two) is treated as a fixed element. That is
    # only a filter: a sample department name repeated on most pages slips through, and so do fields shorter
    # than min_len. What is listed is a candidate to check, not proof of leftover sample text.
    candidates = [p for p, n in seen.items() if not (n >= 2 and n > len(samples) / 2)
                  and len(p.replace(" ", "")) >= min_len]
    found = []
    for idx, slide in enumerate(Presentation(pptx).slides, 1):
        for para in slide_paragraphs(slide):
            for sample in candidates:
                n, shared = longest_common(para, sample)
                if n >= min_len and len(shared.strip()) >= min_len and not shared.strip().isdigit():
                    found.append({"slide": idx, "text": para[:60], "template_text": sample[:60], "shared": shared.strip()})
                    break
    return found


def walk_shapes(shapes, parent=()):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        path = parent + (shape.shape_id,)
        yield path, shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes, path)


def locate_shape(prs, target):
    if not 1 <= target["slide"] <= len(prs.slides) or not target["path"]:
        raise IndexError(f"invalid slide or empty shape path: {target}")
    shapes = prs.slides[target["slide"] - 1].shapes
    for index, ident in enumerate(target["path"]):
        shape = next((s for s in shapes if s.shape_id == ident), None)
        if shape is None:
            raise KeyError(f"shape not found: {target}")
        if index < len(target["path"]) - 1:
            if not hasattr(shape, "shapes"):
                raise KeyError(f"path crosses a non-group shape: {target}")
            shapes = shape.shapes
    return shape


def default_targets(prs):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    targets = []
    for number, slide in enumerate(prs.slides, 1):
        text_parents = set()
        for path, shape in walk_shapes(slide.shapes):
            base = {"slide": number, "path": list(path)}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                targets.append(dict(base, kind="move"))
            if shape.has_text_frame and shape.text_frame.text.strip() and path[:-1] not in text_parents:
                targets.append(dict(base, kind="text"))
                text_parents.add(path[:-1])
            if getattr(shape, "has_table", False):
                targets.append(dict(base, kind="table", cell=[0, 0]))
            if getattr(shape, "has_chart", False):
                targets.append(dict(base, kind="chart"))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                targets.append(dict(base, kind="picture"))
    return targets


def workbook_cells(chart):
    part = chart.part.chart_workbook.xlsx_part
    if part is None:
        raise NotImplementedError("No embedded workbook; external-linked charts are not edited.")
    return workbook_blob_cells(part.blob)


def workbook_blob_cells(blob):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(BytesIO(blob)) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
        if len(sheets) != 1:
            raise NotImplementedError("Only a single embedded worksheet is supported.")
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            for item in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("s:si", ns):
                shared.append("".join(item.itertext()))
        result = {}
        for cell in ET.fromstring(z.read(sheets[0])).findall(".//s:c", ns):
            if cell.find("s:f", ns) is not None:
                raise NotImplementedError("Formula-based workbooks require a dedicated edit test.")
            value = cell.find("s:v", ns)
            if cell.get("t") == "inlineStr":
                result[cell.get("r")] = "".join(cell.find("s:is", ns).itertext())
            elif value is not None and value.text is not None:  # PptxGenJS writes empty <v/> cells
                result[cell.get("r")] = shared[int(value.text)] if cell.get("t") == "s" else float(value.text)
        return result


def chart_snapshot(chart):
    supported = {"BAR_CLUSTERED", "COLUMN_CLUSTERED", "LINE"}
    if chart.chart_type.name not in supported or len(chart.plots) != 1:
        raise NotImplementedError("Automatic editing supports single-plot clustered bar/column and line charts.")
    plot = chart.plots[0]
    if plot.categories.depth != 1:
        raise NotImplementedError("Multi-level categories require a dedicated test.")
    categories = [c.label for c in plot.categories]
    if any(not isinstance(c, str) for c in categories):
        raise NotImplementedError("Date/numeric category axes require a dedicated test.")
    series = [{"name": s.name, "values": list(s.values)} for s in plot.series]
    cells = workbook_cells(chart)
    # replace_data rebuilds the workbook. Only test a known, lossless table layout.
    from pptx.chart.data import CategoryChartData
    data = CategoryChartData()
    data.categories = categories
    for item in series:
        data.add_series(item["name"], item["values"])
    if not equivalent(cells, workbook_blob_cells(data.xlsx_blob)):
        raise NotImplementedError("Nonstandard workbook layout/content requires a dedicated edit test.")
    return {"categories": categories, "series": series, "workbook": cells}


def structure(shape):
    # IDs and endpoint bindings survive a move; this does not prove automatic line routing.
    items = shape.shapes if hasattr(shape, "shapes") else [shape]
    return [{"path": list(path), "type": int(s.shape_type),
             "connections": [dict(el.attrib) for el in s._element.iter()
                             if str(el.tag).split("}")[-1] in ("stCxn", "endCxn")]}
            for path, s in walk_shapes(items)]


def snapshot(shape, target):
    kind = target["kind"]
    if kind == "text":
        return shape.text_frame.text
    if kind == "table":
        return shape.table.cell(*target.get("cell", [0, 0])).text
    if kind == "chart":
        return chart_snapshot(shape.chart)
    if kind == "move":
        return {"left": shape.left, "top": shape.top, "structure": structure(shape)}
    if kind == "picture":
        return [shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom]
    raise ValueError(f"unknown edit kind: {kind}")


def equivalent(actual, expected):
    if isinstance(expected, (int, float)):
        return isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-10)
    if isinstance(expected, dict):
        return isinstance(actual, dict) and actual.keys() == expected.keys() and all(
            equivalent(actual[k], v) for k,v in expected.items())
    if isinstance(expected, list):
        return isinstance(actual, list) and len(actual) == len(expected) and all(
            equivalent(a,b) for a,b in zip(actual, expected))
    return actual == expected


def verify_edits(prs, records):
    for record in records:
        if record["status"] in ("unsupported", "failed"):
            continue
        try:
            record["actual"] = snapshot(locate_shape(prs, record["target"]), record["target"])
            record["status"] = "passed" if equivalent(record["actual"], record["expected"]) else "failed"
        except (KeyError, IndexError, NotImplementedError) as error:
            record["status"], record["detail"] = "failed", str(error)
        except Exception as error:  # report the object instead of aborting the whole check
            record["status"], record["detail"] = "failed", f"checker error {type(error).__name__}: {error}"
    return records


def edit_test(pptx, out_dir, soffice=None, width=1600, targets=None):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.util import Inches
    copy = out_dir / "edit-test.pptx"
    if Path(pptx).resolve() == copy.resolve():
        raise ValueError("Edit-test copy must differ from the source.")
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(pptx)
    selected = default_targets(prs) if targets is None else targets
    identities = [(t["slide"], tuple(t["path"]), t["kind"], tuple(t.get("cell", []))) for t in selected]
    if len(set(identities)) != len(identities):
        raise ValueError("Duplicate edit targets; select each behavior on an object once.")
    records = []
    for number, target in enumerate(selected):
        record = check(f"edit:{number+1}:{target['kind']}", "not_run", target=target)
        records.append(record)
        try:
            shape = locate_shape(prs, target)
            before = snapshot(shape, target)
            record["before"] = before
            kind = target["kind"]
            if kind in ("text", "table"):
                frame = shape.text_frame if kind == "text" else shape.table.cell(*target.get("cell", [0, 0])).text_frame
                runs = [r for p in frame.paragraphs for r in p.runs]
                marker = f" [E{number+1}]"
                if runs:
                    runs[0].text += marker
                else:
                    frame.paragraphs[0].add_run().text = marker
                record["expected"] = frame.text
            elif kind == "chart":
                if not before["series"]:
                    raise NotImplementedError("Empty chart.")
                values = before["series"][0]["values"]
                point = next((i for i,v in enumerate(values) if v is not None and math.isfinite(v)), None)
                if point is None:
                    raise NotImplementedError("No finite numeric point to edit.")
                expected = json.loads(json.dumps(before))
                value = round(values[point] + abs(values[point]) * 0.1 + 1, 6)
                expected["series"][0]["values"][point] = value
                expected["workbook"][f"B{point+2}"] = value
                data = CategoryChartData()
                data.categories = before["categories"]
                for series in expected["series"]:
                    data.add_series(series["name"], series["values"])
                shape.chart.replace_data(data)
                record["expected"] = expected
            elif kind == "move":
                shape.left += Inches(0.1)
                record["expected"] = dict(before, left=before["left"] + Inches(0.1))
            elif kind == "picture":
                shape.crop_left = 0.1 if abs(shape.crop_left-0.1) > 0.001 else 0.05
                record["expected"] = [shape.crop_left, *before[1:]]
                record["detail"] = "Crop round-trip only; replacement and visual crop quality require inspection."
        except NotImplementedError as error:
            record["status"], record["detail"] = "unsupported", str(error)
        except (KeyError, IndexError) as error:
            record["status"], record["detail"] = "failed", str(error)
        except Exception as error:  # one malformed object must not abort the whole report
            record["status"], record["detail"] = "failed", f"checker error {type(error).__name__}: {error}"
    prs.save(copy)
    verify_edits(Presentation(copy), records)
    rendered, render_note = [], "not rendered: LibreOffice or PyMuPDF unavailable"
    if soffice and pymupdf_or_none():
        try:
            rendered, render_note = pptx_to_pngs(copy, out_dir / "edit-test", soffice, width), None
        except (OSError, subprocess.SubprocessError, RuntimeError) as error:
            render_note = f"rendering the edited copy failed: {error}"
    return {"copy": str(copy), "checks": records,
            "selection": "explicit" if targets is not None else "one text per page/group, all tables/charts/groups/pictures",
            "rendered": [str(p) for p in rendered], "render_note": render_note,
            "scope": "Library save/reload; manual edits, auto-connect and data-linked labels need task-specific checks."}


def cmd_pptx(args):
    src = Path(args.input).resolve()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    slides_dir = out / "slides"
    if slides_dir.exists():
        shutil.rmtree(slides_dir)
    # Pages come from LibreOffice, or from --pdf: a PDF exported from this PPTX by WPS, PowerPoint or another
    # renderer. Without either (or without PyMuPDF), rendering and the layout check are reported as not_run;
    # the inventory, font report and edit test still run.
    soffice = find_soffice(args.soffice)
    has_pm = pymupdf_or_none() is not None
    pdf = Path(args.pdf).resolve() if args.pdf else None
    layout, pngs, render_note, render_failed = [], [], None, False
    export_failed = False
    if args.export:
        # The target app itself writes the PDF from the final file: its line breaks and fonts are the ones
        # the audience will see, unlike LibreOffice's approximation.
        try:
            pdf = export_pdf(src, out / f"export-{args.export}.pdf", args.export)
            print(f"exported with {args.export}: {pdf}")
        except (OSError, subprocess.SubprocessError, RuntimeError) as error:
            export_failed = True
            render_note = f"export with {args.export} failed: {error}"
            print("warning: " + render_note, file=sys.stderr)
    zoom_dir = out / "zoom"
    if zoom_dir.exists():
        shutil.rmtree(zoom_dir)
    can_render = has_pm and bool(pdf or soffice) and not export_failed
    if not can_render and not export_failed:
        render_note = ("PyMuPDF not installed" if not has_pm else "LibreOffice not found") + \
            "; rendering and the layout check were not run (inventory, fonts and edit test still run)." + \
            ("" if not has_pm else " Use --export wps|powerpoint, or --pdf with a PDF exported from this PPTX.")
        print("warning: " + render_note, file=sys.stderr)
    elif can_render:
        try:
            zd = zoom_dir if args.zoom else None
            pngs = (pdf_to_pngs(pdf, slides_dir, args.width, layout, zd) if pdf else
                    pptx_to_pngs(src, slides_dir, soffice, args.width, layout, zd))
        except (OSError, subprocess.SubprocessError, RuntimeError) as error:
            render_note, render_failed, layout = f"rendering failed: {error}", True, []
            print("warning: " + render_note, file=sys.stderr)
    renderer = (f"{args.export} export ({EXPORT_APPS[args.export]})" if args.export and pdf else
                f"supplied PDF {pdf.name}" if pdf else "LibreOffice (approximation of PowerPoint)")
    report = {"file": str(src), "renderer": renderer if pngs else None,
              "render_note": render_note, "slides": len(pngs), "fonts": font_inventory(src),
              "inventory": inventory(src), "layout": layout, "text": pptx_text(src, args.visual_first)}
    frame = deck_frame([(item["slide"], item.get("frame")) for item in layout], "pt", tolerance=2, threshold=12)
    report["deck_layout"] = [frame] if frame else []
    if args.template:
        # Sample text of the company template left in the deck (department names, dates, sample titles).
        report["template_text"] = template_sample_text(src, Path(args.template).resolve())
    if pdf:
        # The checks only cover the version that was exported: record both files so they can be matched.
        stale = pdf.stat().st_mtime < src.stat().st_mtime
        report["pdf"] = {"file": str(pdf), "sha256": sha256(pdf), "pptx_sha256": sha256(src),
                         "older_than_pptx": stale}
        if stale:
            warning = ("the PDF is older than the PPTX: export it again from the final file, "
                       "or confirm it matches before relying on these checks.")
            render_note = report["render_note"] = f"{render_note} {warning}" if render_note else warning
            print("warning: " + warning, file=sys.stderr)
    layout_dir = out / "layout"
    if layout_dir.exists():
        shutil.rmtree(layout_dir)
    for item in layout:
        if item["findings"]:
            draw_layout(slides_dir / f"slide-{item['slide']:02d}.png", item["findings"],
                        layout_dir / f"slide-{item['slide']:02d}.png", item["scale"])
    layout_errors = sum(f["severity"] == "error" for item in layout for f in item["findings"])
    layout_warnings = sum(f["severity"] == "warning" for item in layout for f in item["findings"]) + len(report["deck_layout"])
    if args.edit_test:
        if report["inventory"] is None:
            report["edit_test"] = "skipped: python-pptx not installed"
        else:
            targets = json.loads(Path(args.edit_targets).read_text(encoding="utf-8")) if args.edit_targets else None
            report["edit_test"] = edit_test(src, out, soffice if soffice and has_pm else None, args.width, targets)
    if pngs:
        make_sheet(pngs, out / "sheet.png")
        print(f"slides: {len(pngs)} -> {slides_dir}")
        print(f"sheet  -> {out / 'sheet.png'} (+ -gray, -blur)")
    else:
        print(f"slides: not rendered ({render_note})")
    print(f"fonts: {report['fonts']}")
    if report["inventory"]:
        for c in report["inventory"]:
            print("  " + json.dumps(c, ensure_ascii=False))
    if "edit_test" in report:
        edit = report["edit_test"]
        if isinstance(edit, dict):
            # Name what was actually tested, so a record cannot claim object types the deck does not have.
            kinds = {}
            for c in edit["checks"]:
                kind = (c.get("target") or {}).get("kind", "?")
                kinds[kind] = kinds.get(kind, 0) + 1
            inv = report["inventory"] or []
            absent = [name for name, key in (("chart", "charts"), ("table", "tables"))
                      if not any(s[key] for s in inv)]
            edit["targets_by_kind"] = kinds
            print(f"edit test: {checks_status(edit['checks'])}, {len(edit['checks'])} targets ("
                  + ", ".join(f"{k} {n}" for k, n in kinds.items()) + ")"
                  + (f"; this deck has no native {' or '.join(absent)}" if absent else "")
                  + "; details in report.json")
        else:
            print(f"edit test: {edit}")
    if pngs:
        print(f"layout: {layout_errors} errors, {layout_warnings} warnings"
              + (f" -> annotated pages in {layout_dir}" if layout_errors or layout_warnings else ""))
    else:
        print("layout: not run (needs a rendering)")
    for item in layout:
        for f in item["findings"]:
            print(f"  slide {item['slide']:02d} {f['severity']:<7} {f['kind']:<18} “{f['text']}” {f['detail']}")
    for f in report["deck_layout"]:
        print(f"  deck     {f['severity']:<7} {f['kind']:<18} {f['detail']}")
    if "template_text" in report:
        print(f"text shared with the template's sample slides: {len(report['template_text'])} candidates to check "
              "(keep fixed brand text and text the material needs; replace sample text; list what you cannot "
              "confirm as 待确认). Text on most template slides is skipped, so repeated sample names and fields "
              "under 5 characters can slip through: look at the cover, dates and titles yourself")
        for f in report["template_text"]:
            print(f"  slide {f['slide']:02d} warning template_text       “{f['text']}” shares “{f['shared']}” with "
                  f"the template's “{f['template_text']}”")
    if args.zoom and pngs:
        print(f"zoom   -> {zoom_dir} (four enlarged quarters per slide: look again at text, breaks and edges)")
    if report["text"]:
        print_text(report["text"])

    expected_pages = len(report["inventory"]) if report["inventory"] is not None else None
    if export_failed:
        render_status = "failed"
    elif not can_render:
        render_status = "not_run"
    elif render_failed:
        render_status = "failed"
    else:
        render_status = ("unsupported" if expected_pages is None else
                         "passed" if len(pngs) == expected_pages else "failed")
    report["checks"] = [check("render", render_status, render_note or "Compare rendered pages with source slides.",
                              expected=expected_pages, actual=len(pngs)),
                        check("inventory", "passed" if report["inventory"] is not None else "unsupported")]
    if pngs:
        report["checks"].append(check(
            "layout", "failed" if layout_errors else "passed",
            f"{layout_errors} errors (outside slide, overlapping text, text hidden under a shape), {layout_warnings} "
            "warnings (orphan line, text past its box, small font, low contrast, placeholder, line through text, "
            "off-centre content frame). "
            + (f"Measured on the PDF exported by {args.export}. " if args.export and pdf else
               "Measured on the supplied PDF; it covers only the version that was exported. " if pdf else
               "Measured on the LibreOffice rendering; PowerPoint/WPS text metrics can differ. ")
            + "Contrast is read from the rendering; text on photos is not measured.",
            errors=layout_errors, warnings=layout_warnings))
    else:
        report["checks"].append(check("layout", "not_run", "Needs a rendering; " + (render_note or "")))
    edit = report.get("edit_test")
    if isinstance(edit, dict) and edit["checks"]:
        report["checks"].extend(edit["checks"])
    else:
        report["checks"].append(check("edit_test", "not_run" if edit is None or isinstance(edit, dict) else "unsupported",
                                      "No edit targets verified. Supply --edit-test and task-specific targets if needed."))
    write_report(report, out, args.strict)


# ---------------------------------------------------------------- CLI

VISUAL_FIRST_HELP = (f"the deck is visual-first: hint at any paragraph or box of points over {LONG_TEXT} 字, "
                     f"any slide over {PAGE_TEXT} 字 of sentences (lines over {LABEL_TEXT} 字; short labels "
                     "are not counted; a Latin word or number counts as 2) and any slide with "
                     f"{BARE_TEXT} 字 or more but no picture, illustration, chart or table; draw the information "
                     "instead of deleting it; hints only, never a failed check")
ZOOM_HELP = "also write four enlarged, overlapping quarters of every slide to zoom/ for a second, closer look"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    h = sub.add_parser("html", help="screenshot and check an HTML deck")
    h.add_argument("input")
    h.add_argument("--out", required=True)
    h.add_argument("--scale", type=float, default=1.25, help="device scale factor (1.25 -> 1600px wide)")
    h.add_argument("--browser", choices=["auto", "chromium", "chrome", "msedge"], default="auto")
    h.add_argument("--no-pdf", action="store_true", help="skip the print-to-PDF check")
    h.add_argument("--strict", action="store_true", help="nonzero exit unless required automatic checks pass")
    h.add_argument("--visual-first", action="store_true", help=VISUAL_FIRST_HELP)
    h.add_argument("--zoom", action="store_true", help=ZOOM_HELP)
    h.set_defaults(func=cmd_html)

    x = sub.add_parser("pptx", help="render a PPTX through LibreOffice, WPS or PowerPoint (or check a PDF "
                                    "exported from it) and inspect native objects")
    x.add_argument("input")
    x.add_argument("--out", required=True)
    x.add_argument("--width", type=int, default=1600)
    x.add_argument("--soffice")
    source = x.add_mutually_exclusive_group()
    source.add_argument("--export", choices=sorted(EXPORT_APPS),
                        help="Windows: have WPS or PowerPoint open the final PPTX read-only and export the PDF that "
                             "the pages and the layout check use (the app's own line breaks and fonts)")
    source.add_argument("--pdf", help="use this PDF, exported from the same final PPTX by WPS, PowerPoint or another "
                                      "renderer, for the pages and the layout check instead of LibreOffice")
    x.add_argument("--edit-test", action="store_true", help="edit text/table/chart in a copy, reload and render")
    x.add_argument("--edit-targets", help="JSON list of slide (1-based), shape-id path, kind, optional table cell")
    x.add_argument("--strict", action="store_true", help="nonzero exit unless required automatic checks pass")
    x.add_argument("--visual-first", action="store_true", help=VISUAL_FIRST_HELP)
    x.add_argument("--zoom", action="store_true", help=ZOOM_HELP)
    x.add_argument("--template", help="the company template the deck was built from: list text the deck still "
                                      "shares with the template's sample slides (department, date, sample titles)")
    x.set_defaults(func=cmd_pptx)

    s = sub.add_parser("sheet", help="contact sheet from PNG files or a directory")
    s.add_argument("inputs", nargs="+")
    s.add_argument("--out", required=True)
    s.add_argument("--cols", type=int)
    s.add_argument("--thumb", type=int, default=400)
    s.set_defaults(func=lambda a: print(make_sheet(
        [q for i in a.inputs for q in (sorted(Path(i).glob("*.png")) if Path(i).is_dir() else [Path(i)])],
        a.out, a.cols, a.thumb)))

    c = sub.add_parser("compare", help="place images side by side")
    c.add_argument("inputs", nargs="+")
    c.add_argument("--out", required=True)
    c.add_argument("--labels", help="comma-separated labels (default A, B, C ...)")
    c.add_argument("--height", type=int, default=540)
    c.set_defaults(func=lambda a: print(make_compare(
        a.inputs, a.out, a.labels.split(",") if a.labels else None, a.height)))

    args = ap.parse_args()
    if getattr(args, "edit_targets", None) and not args.edit_test:
        ap.error("--edit-targets requires --edit-test")
    if getattr(args, "pdf", None) and not Path(args.pdf).is_file():
        ap.error(f"--pdf: no such file {args.pdf}")
    if getattr(args, "template", None) and not Path(args.template).is_file():
        ap.error(f"--template: no such file {args.template}")
    args.func(args)


if __name__ == "__main__":
    main()
