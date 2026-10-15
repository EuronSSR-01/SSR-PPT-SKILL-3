#!/usr/bin/env python3
"""Which characters a font cannot show: check a font file against a deck, a page or a piece of text.

  font_coverage.py fonts/SmileySans-Oblique.ttf --pptx build/deck.pptx
  font_coverage.py C:/Windows/Fonts/msyh.ttc#0 fonts/ZCOOLXiaoWei-Regular.ttf --text "珺堃喆 · 〇 ℃"
  font_coverage.py fonts/LXGWWenKai-Regular.ttf --html build/index.html --file planning.md

For each font it prints the family name, how many of the 6,763 GB2312 Chinese characters it has, and the
characters of your text it lacks. A missing character is drawn by some other font (PowerPoint and WPS pick
one per character, often a Song or DengXian face), so a line in a display font suddenly has one plain
character in it. Fix by choosing a font with more characters for that text, rewording, or giving those
runs an explicit fallback font. TTC collections: add #index (default #0). Needs fontTools; --pptx needs
python-pptx.
"""
import argparse
import re
import sys
from pathlib import Path


def gb2312_hanzi():
    out = []
    for b1 in range(0xB0, 0xF8):
        for b2 in range(0xA1, 0xFF):
            try:
                ch = bytes([b1, b2]).decode("gb2312")
            except UnicodeDecodeError:
                continue
            if "\u4e00" <= ch <= "\u9fff":
                out.append(ch)
    return out


def pptx_text(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # group
                walk(sh.shapes)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                chart = sh.chart
                if chart.plots:
                    parts.extend(str(c) for c in chart.plots[0].categories)
    for slide in prs.slides:
        walk(slide.shapes)
    return "\n".join(parts)


def html_text(path):
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"<(script|style)\b.*?</\1>", " ", text, flags=re.S | re.I)
    return re.sub(r"<[^>]+>", " ", text)


def load_font(spec):
    from fontTools.ttLib import TTFont
    path, _, index = spec.partition("#")
    font = TTFont(path, fontNumber=int(index or 0), lazy=True)
    try:
        name = font["name"].getDebugName(1) or Path(path).stem
        cmap = dict(font.getBestCmap() or {})
    finally:
        font.close()
    return name, cmap


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("fonts", nargs="+", help="font files (TTF/OTF, or TTC with #index)")
    ap.add_argument("--pptx", action="append", default=[])
    ap.add_argument("--html", action="append", default=[])
    ap.add_argument("--file", action="append", default=[], help="plain text or Markdown files")
    ap.add_argument("--text", action="append", default=[])
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    text = "".join(args.text)
    text += "".join(pptx_text(p) for p in args.pptx)
    text += "".join(html_text(p) for p in args.html)
    text += "".join(Path(p).read_text(encoding="utf-8", errors="replace") for p in args.file)
    wanted = sorted({c for c in text if not c.isspace() and ord(c) > 0x20})
    common = gb2312_hanzi()
    worst = 0
    for spec in args.fonts:
        name, cmap = load_font(spec)
        have = sum(1 for c in common if ord(c) in cmap)
        missing = [c for c in wanted if ord(c) not in cmap]
        worst = max(worst, len(missing))
        print(f"{name}  ({spec})")
        print(f"  GB2312 常用汉字 {have}/{len(common)}")
        if wanted:
            print(f"  本文缺 {len(missing)} 个字符：{''.join(missing) if missing else '—'}")
    sys.exit(1 if worst else 0)


if __name__ == "__main__":
    main()
